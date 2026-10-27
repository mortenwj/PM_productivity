#!/usr/bin/env python3
"""Fill Enablement product sheet template (75122) for 360° Flow module — P00285.

Source of truth for text: 360_Flow_Product_Sheet_ENG.md (keep in sync with manual PPT edits).
"""

from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent
TEMPLATE = (
    ROOT
    / "reference-product-sheets"
    / "75122 360_ProductName_LAN 617079_3_0.pptx"
)
OUTPUT = ROOT / "360_Flow_Product_Sheet_ENG.pptx"

PRODUCT_NUMBER = "P00285"

TITLE = f"360° Flow\vProduct number: {PRODUCT_NUMBER}\v"

SHORT_DESCRIPTION = (
    "Governed, process-oriented case handling in Public 360° — from registration "
    "to completion on the case."
)

DESCRIPTION = """Many public-sector organizations still coordinate case work through tasks, documents, and informal checklists. Work is hard to follow end to end, deadlines are difficult to manage, and managers lack a consistent view of status and workload.

360° Flow is a Public 360° capability for process-oriented case processing. Work on a case is organized through a process model with phases, checklists, milestones, and process activities — with clear status, deadlines, and reporting on the case. Caseworkers work in phase context; managers can follow progress through worklists and portfolio views.

Deliver consistent case handling, better overview, and predictable progression — with standard platform security, auditability, and role enforcement, without adding a separate process engine outside Public 360°."""

ADVANTAGES_AND_REQUIREMENTS = """Advantages:
Efficiency — structure work on the case Process models and phases give caseworkers a clear path from start to finish;
Quality — checklists and phase context improve overview and help catch gaps before they become errors.
Equal treatment — consistent steps per case. Milestones and process activities support the same progression for similar cases,
Predictability — deadlines and status you can trust
Usability — navigate and prioritise at scale ProcessView and portfolio-oriented worklists

Technical requirements:
Public 360° version 6.12 or newer (online, domestic and locally installed)"""


def set_shape_text(shape, text: str) -> None:
    if not shape.has_text_frame:
        return
    shape.text_frame.text = text


def find_shape_by_name(slide, name: str):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    return None


def find_advantages_shape(slide):
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        if "Advantages:" in shape.text_frame.text or "Technical requirements:" in shape.text_frame.text:
            return shape
    return None


def main() -> None:
    if not TEMPLATE.is_file():
        raise SystemExit(f"Template not found: {TEMPLATE}")

    shutil.copy2(TEMPLATE, OUTPUT)
    prs = Presentation(str(OUTPUT))
    slide = prs.slides[0]

    title_shape = find_shape_by_name(slide, "Title 1")
    if title_shape is None:
        raise SystemExit("Title 1 shape not found")
    set_shape_text(title_shape, TITLE)

    desc_shape = find_shape_by_name(slide, "Text Placeholder 4")
    if desc_shape is None:
        raise SystemExit("Text Placeholder 4 not found")
    set_shape_text(desc_shape, DESCRIPTION)

    short_shape = find_shape_by_name(slide, "Text Placeholder 7")
    if short_shape is None:
        raise SystemExit("Text Placeholder 7 not found")
    set_shape_text(short_shape, SHORT_DESCRIPTION)

    adv_shape = find_advantages_shape(slide)
    if adv_shape is None:
        raise SystemExit("Advantages / requirements shape not found")
    set_shape_text(adv_shape, ADVANTAGES_AND_REQUIREMENTS)

    prs.core_properties.title = "360° Flow"
    prs.core_properties.subject = "360 product sheet ENG"
    prs.core_properties.keywords = PRODUCT_NUMBER

    prs.save(str(OUTPUT))
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
