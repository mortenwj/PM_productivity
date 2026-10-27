#!/usr/bin/env python3
"""Fill Enablement product sheet template (75122) for 360° Flow implementation package."""

from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent
TEMPLATE = (
    ROOT
    / "reference-product-sheets"
    / "75122 360_ProductName_LAN 617079_3_0.pptx"
)
OUTPUT = ROOT / "360_Flow_Implementation_Package_Product_Sheet_ENG.pptx"

TITLE = (
    "360° Flow — Implementation Package (2 Processes)\v"
    "Product number: [TBD — Enablement catalogue / FLOW-IMPL-2PROC]\v"
)

SHORT_DESCRIPTION = (
    "Operational case progression in Public 360° — two processes delivered as a "
    "predictable, fixed-scope package."
)

DESCRIPTION = """Many public-sector organizations handle cases through fragmented tasks, documents, and manual coordination. That leads to inconsistent service, missed deadlines, and weak management oversight. Custom development is costly and hard to maintain through platform upgrades.

360° Flow brings governed, end-to-end process support into Public 360° — process models, phases, checklists, milestones, and ProcessView on the case — without standalone BPM or bespoke development.

This implementation package is a fixed-scope delivery engagement: Tieto designs, configures, tests, and go-lives two production-ready processes built from standard shipped building blocks. When you sign the delivery contract, platform capacity for two active process models is included at no additional platform fee under the introductory pilot offer (delivery contracts signed by 31 December 2026, subject to change).

The package includes process design through go-live for two processes, project management, configuration documentation, handover orientation, and platform entitlement for two active process models bundled with the signed delivery contract. A third process requires separate implementation services and capacity fees."""

ADVANTAGES_AND_REQUIREMENTS = """Advantages:
Predictable delivery — fixed scope and milestones, not open-ended consulting.
Faster time to value — proven building blocks and repeatable delivery patterns.
Operational clarity — phases, checklists, and milestones on every case.
Compliance-ready foundation — progression auditability and role enforcement on the platform.
Included platform capacity — two active process models at no extra platform fee when bundled with this delivery contract (pilot offer 2026).

Technical requirements:
Public 360° subscription (version 6.10 or newer)
360° Flow pilot enablement on the tenant
Signed delivery contract for this implementation package"""


def set_shape_text(shape, text: str) -> None:
    if not shape.has_text_frame:
        return
    shape.text_frame.text = text


def find_shape_by_name(slide, name: str):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    return None


def find_advantages_shape(slide):
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        if "Advantages:" in shape.text_frame.text or "Technical requirements:" in shape.text_frame.text:
            return shape
    return None


def main() -> None:
    if not TEMPLATE.is_file():
        raise SystemExit(f"Template not found: {TEMPLATE}")

    shutil.copy2(TEMPLATE, OUTPUT)
    prs = Presentation(str(OUTPUT))
    slide = prs.slides[0]

    title_shape = find_shape_by_name(slide, "Title 1")
    if title_shape is None:
        raise SystemExit("Title 1 shape not found")
    set_shape_text(title_shape, TITLE)

    desc_shape = find_shape_by_name(slide, "Text Placeholder 4")
    if desc_shape is None:
        raise SystemExit("Text Placeholder 4 not found")
    set_shape_text(desc_shape, DESCRIPTION)

    short_shape = find_shape_by_name(slide, "Text Placeholder 7")
    if short_shape is None:
        raise SystemExit("Text Placeholder 7 not found")
    set_shape_text(short_shape, SHORT_DESCRIPTION)

    adv_shape = find_advantages_shape(slide)
    if adv_shape is None:
        raise SystemExit("Advantages / requirements shape not found")
    set_shape_text(adv_shape, ADVANTAGES_AND_REQUIREMENTS)

    prs.core_properties.title = "360° Flow — Implementation Package (2 Processes)"
    prs.core_properties.subject = "360 product sheet ENG"
    prs.core_properties.keywords = "FLOW-IMPL-2PROC"

    prs.save(str(OUTPUT))
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
