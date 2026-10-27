"""Update Flow training videos catalogue spreadsheet."""
from __future__ import annotations

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PATH = (
    r"C:\Users\jensemor\OneDrive - Tietoevry"
    r"\Technology & Development - Public 360 - ProcessModel"
    r"\Training\training videos.xlsx"
)

HEADERS = ["Name", "Description", "Audience", "Priority", "Status", "Owner"]

ROWS = [
    (
        "Flow introduction",
        "Introduction to 360° Flow (formerly PAF). Process models, progress plans, ProcessView, "
        "milestones, activities. Pilot/internal scope; what to watch next in the series.",
        "All",
        1,
        "New",
        None,
    ),
    (
        "What Flow is — and is not",
        "Governed operational execution layer inside Public 360° — not generic BPM or low-code. "
        "Hosted Processes as future commercial layer on top of Flow.",
        "All / Leadership",
        1,
        "New",
        None,
    ),
    (
        "Process Models",
        "What is a process model? Lifecycle structure: stages, milestones, activities, binding "
        "to cases. Foundation for all other Flow topics.",
        "All",
        1,
        "New",
        None,
    ),
    (
        "Progress plans vs process models",
        "Canonical distinction: process model = structure; progress plan = executable flow on the "
        "case. Binding a model, starting progression, subprocesses.",
        "All",
        1,
        "New",
        None,
    ),
    (
        "ProcessView: your operational home on the case",
        "Tour of ProcessView (DMS/CaseDetails/ProcessView): phases, activity list, "
        "execute/edit/cancel/reactivate, dialogs, where process status lives on the case.",
        "Case workers",
        1,
        "New",
        None,
    ),
    (
        "Activities and the activity dialog",
        "What an activity is; responsible party and due dates; run-without-dialog; complete, "
        "repeat, or cancel steps and what happens in the UI.",
        "Case workers",
        1,
        "New",
        None,
    ),
    (
        "Working with processes",
        "Case-handler execution focus: progressing the case through configured processes and "
        "steps. Complements ProcessView tour (navigation and columns).",
        "Case workers",
        1,
        "New",
        None,
    ),
    (
        "Stages and stage templates",
        "Stages in the process model lifecycle; stage templates; reuse and configuration "
        "patterns for process builders.",
        "Process Builders",
        1,
        "New",
        None,
    ),
    (
        "Milestones and triggers",
        "Introduction to milestones: setup, triggers from processes, tasks, and documents "
        "(including dispatch/receive). How milestones support reporting and visibility.",
        "Process Builders",
        1,
        "New",
        None,
    ),
    (
        "Document-linked milestones",
        "Milestones triggered from document dispatch/receive and related orchestration. What "
        "pilot tenants should validate (categories, templates).",
        "Process Builders",
        1,
        "New",
        None,
    ),
    (
        "Stage defined processes and building blocks",
        "Using predefined stage processes. How to combine processes, arguments, and governed "
        "Flow building blocks from the process library.",
        "Process Builders",
        1,
        "New",
        None,
    ),
    (
        "Setting up your first process",
        "Creating a process in admin: stage templates, process library, configuring arguments. "
        "For process builders with implementation support.",
        "Process Builders / Implementation",
        1,
        "New",
        None,
    ),
    (
        "Enabling Flow in your tenant",
        "Pilot enablement: implementation-led activation, AsyncProcessModelLoading, relation "
        "to FT:ProcessManagementModule, who to contact.",
        "Admins / Implementation",
        1,
        "New",
        None,
    ),
    (
        "Flow roles: who does what",
        "Case worker, registrar, process owner, manager. Registrar-gated PMC; stage permissions "
        "and pilot ACL considerations.",
        "All",
        2,
        "New",
        None,
    ),
    (
        "DMN rules in Flow (introduction)",
        "What rules do at runtime; when DMN applies; what builders configure vs what case workers "
        "see. Not a full DMN authoring course.",
        "Process Builders",
        2,
        "New",
        None,
    ),
    (
        "Notifications and email from processes",
        "How configured notifications fire; email patterns; what case workers receive vs what "
        "admins configure. Note transitional FFUI/legacy surfaces where relevant.",
        "Process Builders / Case workers",
        2,
        "New",
        None,
    ),
    (
        "Document flows: approval, review, dispatch",
        "How approval, review, and dispatch appear in the progress plan. Mixed legacy/React "
        "routing; pilot caveats.",
        "Process Builders",
        2,
        "New",
        None,
    ),
    (
        "Sub-processes and process dialogs",
        "When subprocesses open; nested execution; returning to the parent process.",
        "Case workers / Process Builders",
        2,
        "New",
        None,
    ),
    (
        "Wait tasks and templates",
        "What is a wait task, and how it differs from milestones. Connecting expired and "
        "completed actions. Templates with wait-task data.",
        "Process Builders",
        2,
        "New",
        None,
    ),
    (
        "Share with (pilot overview)",
        "What Share With is for; toolbox vs ProcessView gap. Tenant verification checklist. "
        "Emerging — label pilot content.",
        "Process Builders",
        2,
        "New",
        None,
    ),
    (
        "PMC: import and export process models",
        "Admin UI for moving models between environments; restrictions on stage templates. "
        "Emerging — not full self-service for all tenants.",
        "Process Builders / Admins",
        2,
        "New",
        None,
    ),
    (
        "Pattern: consultation and approval building blocks",
        "Shorter pattern walkthrough (10–15 min) before the full deficiency-letter capstone.",
        "Process Builders",
        2,
        "New",
        None,
    ),
    (
        "Portfolio management (pilot / emerging)",
        "Desktops, webparts, and worklists direction; what works today vs roadmap. Emerging — "
        "ADO #411634 / #414651; not GA.",
        "Managers / Case workers",
        2,
        "New",
        None,
    ),
    (
        "Automation example: deficiency letter",
        "Capstone walkthrough: automation steps, templates, wait tasks, dispatch. Watch after "
        "wait tasks and document-linked milestones.",
        "Process Builders",
        2,
        "New",
        None,
    ),
    (
        "Localization in multilingual pilots",
        "Configure dialogs (shipped) vs ProcessView lists/dropdowns (emerging). For pilot "
        "tenants with multilingual requirements.",
        "Process Builders",
        3,
        "New",
        None,
    ),
    (
        "Hosted Processes and sector packages (preview)",
        "Commercial/service layer on top of Flow. Out of pilot scope unless explicitly enabled "
        "for the tenant.",
        "Leadership / Sales / PM",
        3,
        "New",
        None,
    ),
]


def main() -> None:
    wb = load_workbook(PATH)
    ws = wb.active
    ws.title = "Catalogue"

    ws.delete_rows(1, ws.max_row)

    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row_idx, row in enumerate(ROWS, 2):
        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for col_idx, width in enumerate([42, 72, 28, 10, 12, 18], 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{len(ROWS) + 1}"

    wb.save(PATH)
    print(f"Updated {PATH}")
    print(f"  Sheet: {ws.title}")
    print(f"  Videos: {len(ROWS)}")
    print(f"  Header fix: Prority -> Priority")


if __name__ == "__main__":
    main()
