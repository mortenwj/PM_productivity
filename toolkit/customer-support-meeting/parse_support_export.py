#!/usr/bin/env python3
"""Parse ServiceNow/CIM support case Excel export for customer meeting prep."""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

CS_PATTERN = re.compile(r"\bCS\d+\b", re.IGNORECASE)

COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "case_number": ("case", "case number", "number", "saksnummer", "cs number", "support number"),
    "description": ("description", "beskrivelse"),
    "comments": ("comments", "comment", "kommentarer", "kommentar"),
    "team_main": ("pm_group_main", "pm group main", "team", "ansvarlig team"),
    "team_module": ("pm_group_module", "pm group module", "module", "modul"),
    "status": ("status", "state", "sakstatus"),
    "ai_summary": ("ai summary", "ai oppsummering", "ai-oppsummering", "summary ai"),
}

LETTER_FALLBACK: dict[str, str] = {
    "description": "F",
    "comments": "G",
    "team_main": "H",
    "ai_summary": "J",
}


def col_index(letter: str) -> int:
    idx = 0
    for char in letter.upper():
        idx = idx * 26 + (ord(char) - ord("A") + 1)
    return idx - 1


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().lower().split())


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def find_header_row(rows: list[tuple[Any, ...]], max_scan: int = 15) -> int:
    best_row = 0
    best_score = -1
    for i, row in enumerate(rows[:max_scan]):
        headers = [normalize_header(c) for c in row]
        score = sum(
            1
            for aliases in COLUMN_ALIASES.values()
            for alias in aliases
            if any(alias in h or h in alias for h in headers if h)
        )
        if score > best_score:
            best_score = score
            best_row = i
    return best_row


def map_columns(headers: list[Any]) -> dict[str, int | None]:
    normalized = [normalize_header(h) for h in headers]
    mapping: dict[str, int | None] = {key: None for key in COLUMN_ALIASES}
    mapping["case_number"] = None

    for key, aliases in COLUMN_ALIASES.items():
        for idx, header in enumerate(normalized):
            if not header:
                continue
            if any(alias == header or alias in header or header in alias for alias in aliases):
                mapping[key] = idx
                break

    for key, letter in LETTER_FALLBACK.items():
        if mapping.get(key) is None:
            mapping[key] = col_index(letter)

    if mapping["case_number"] is None:
        for idx, header in enumerate(normalized):
            if "cs" in header or "case" in header or "saks" in header:
                mapping["case_number"] = idx
                break

    return mapping


def infer_case_number(row: tuple[Any, ...], mapping: dict[str, int | None]) -> str:
    idx = mapping.get("case_number")
    if idx is not None and idx < len(row):
        text = cell_text(row[idx])
        match = CS_PATTERN.search(text)
        if match:
            return match.group(0).upper()
        if text:
            return text

    for value in row:
        match = CS_PATTERN.search(cell_text(value))
        if match:
            return match.group(0).upper()
    return ""


def get_field(row: tuple[Any, ...], mapping: dict[str, int | None], key: str) -> str:
    idx = mapping.get(key)
    if idx is None or idx >= len(row):
        return ""
    return cell_text(row[idx])


def is_open_status(status: str) -> bool:
    normalized = status.strip().lower()
    if not normalized:
        return True
    closed_markers = (
        "closed",
        "close",
        "resolved",
        "cancelled",
        "canceled",
        "lukket",
        "avsluttet",
        "løst",
    )
    return not any(marker in normalized for marker in closed_markers)


def parse_workbook(path: Path, team_filter: list[str] | None = None) -> dict[str, Any]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        raise ValueError("Workbook is empty")

    header_row_idx = find_header_row(rows)
    headers = list(rows[header_row_idx])
    mapping = map_columns(headers)
    data_rows = rows[header_row_idx + 1 :]

    cases: list[dict[str, Any]] = []
    for row in data_rows:
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        case_number = infer_case_number(row, mapping)
        team_main = get_field(row, mapping, "team_main")
        status = get_field(row, mapping, "status")

        if team_filter and team_main not in team_filter:
            continue

        cases.append(
            {
                "case_number": case_number,
                "team_main": team_main,
                "team_module": get_field(row, mapping, "team_module"),
                "status": status,
                "is_open": is_open_status(status),
                "description": get_field(row, mapping, "description"),
                "comments": get_field(row, mapping, "comments"),
                "ai_summary": get_field(row, mapping, "ai_summary"),
            }
        )

    team_counts = Counter(c["team_main"] or "(blank)" for c in cases)
    open_by_team: dict[str, int] = defaultdict(int)
    closed_by_team: dict[str, int] = defaultdict(int)
    modules_open: dict[str, Counter[str]] = defaultdict(Counter)

    for case in cases:
        team = case["team_main"] or "(blank)"
        if case["is_open"]:
            open_by_team[team] += 1
            modules_open[team][case["team_module"] or "(blank)"] += 1
        else:
            closed_by_team[team] += 1

    return {
        "source_file": str(path),
        "parsed_at": datetime.now().isoformat(timespec="seconds"),
        "sheet_name": ws.title,
        "header_row": header_row_idx + 1,
        "total_rows": len(data_rows),
        "case_count": len(cases),
        "columns_detected": {
            key: (headers[idx] if idx is not None and idx < len(headers) else letter)
            for key, idx in mapping.items()
            for letter in [LETTER_FALLBACK.get(key, "")]
            if key != "case_number" or idx is not None
        },
        "team_filter": team_filter or [],
        "teams": [
            {
                "team_main": team,
                "total": count,
                "open": open_by_team.get(team, 0),
                "closed": closed_by_team.get(team, 0),
                "open_modules": dict(modules_open.get(team, {})),
            }
            for team, count in team_counts.most_common()
        ],
        "cases": cases,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path, help="Path to Excel export (.xlsx)")
    parser.add_argument(
        "--teams",
        nargs="*",
        default=None,
        help="Filter to PM_group_main values (space-separated)",
    )
    parser.add_argument(
        "--output-json",
        type=Path,
        default=Path("outputs/customer-support/_parsed.json"),
        help="Where to write parsed JSON",
    )
    args = parser.parse_args()

    if not args.input.exists():
        raise SystemExit(f"Input file not found: {args.input}")

    result = parse_workbook(args.input, team_filter=args.teams)
    args.output_json.parent.mkdir(parents=True, exist_ok=True)
    args.output_json.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"Parsed {result['case_count']} cases from {args.input}")
    print(f"Teams: {len(result['teams'])}")
    for team in result["teams"][:10]:
        print(
            f"  - {team['team_main']}: {team['total']} total "
            f"({team['open']} open, {team['closed']} closed)"
        )
    print(f"JSON written to {args.output_json}")


if __name__ == "__main__":
    main()
