"""Load module revenue from External Revenue Excel (openpyxl; no pandas)."""

from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

# Product name (Excel) -> ImpactedModule / wiki module name.
# Keep aligned with toolkit/scripts/generate_revenue_leaderboard.py
PRODUCT_TO_MODULE: dict[str, str] = {
    "WebCruiter": "WebCruiter",
    "eArchive Historic search": "eArchive",
    "360° eEiendom": "eEiendom",
    "Consulting Customized": "Consulting",
    "Business 360°": "Business 360°",
    "360° Online Public": "Public 360°",
    "360° Scanning": "Scanning",
    "360° Integrasjon mot Folkeregisterert (NO)": "Integrations",
    "360° Board and Council": "Board meeting",
    "360° OpenGov Cases and Documents": "Public link",
    "360° eSigning Connect": "eSignature",
    "360° Contract Management": "Contract Management",
    "360° File Viewer": "Fileviewer",
    "360° eManager": "eManager",
    "Local hosting": "Local hosting",
    "360° Restricted mode": "Restricted Mode",
    "360° Integration with Sharepoint": "Integrations",
    "360° Mimir": "Mimir",
    "360° eServices Connect": "eServices",
    "360° eLearning": "eLearning",
    "eDoc Consulting": "Consulting",
    "P360° Lead offerings": "Public 360°",
    "360° Dispatch": "Dispatch",
    "360° eWorker": "eWorker",
    "360° AD Sync Service": "Usersync",
    "360° Format Converter": "Format converter",
    "Public 360°": "Public 360°",
    "360° PDF Signing": "Personal Digital Signing",
    "360° Archive Format Validator": "File validation",
    "360° Import Center": "File center",
    "360° Process Engine": "Process Engine",
    "360° AutoSaver": "Document",
    "Plan & Build 360°": "Plan&Build",
    "360° Collect": "Collect",
}


def _pkg_dir() -> Path:
    return Path(__file__).resolve().parent


def load_taxonomy() -> dict[str, Any]:
    path = _pkg_dir() / "module_taxonomy.json"
    return json.loads(path.read_text(encoding="utf-8"))


def load_scoring_config() -> dict[str, Any]:
    path = _pkg_dir() / "scoring_config.json"
    return json.loads(path.read_text(encoding="utf-8"))


def extract_product_name(full_name: str) -> str:
    if not isinstance(full_name, str):
        return ""
    parts = full_name.strip().split(None, 1)
    return parts[1] if len(parts) > 1 else full_name


def _to_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _read_export_sheet(excel_path: Path) -> tuple[list[list[Any]], str | None]:
    """Return data rows (from row index 2) and optional growth column name."""
    wb = load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb["Export"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if len(rows) < 3:
        return [], None

    header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    last_header = header[-1] if header else ""
    data_rows = [list(r) for r in rows[2:]]

    if "total" in last_header or len(rows[0]) > 10:
        return data_rows, None

    # External Revenue format — growth in column index 8 if present
    growth_col = "Delta YTD PY" if len(rows[0]) > 8 else None
    return data_rows, growth_col


def resolve_module(product_name: str) -> str | None:
    module = PRODUCT_TO_MODULE.get(product_name)
    if module:
        return module
    for key, mod in PRODUCT_TO_MODULE.items():
        if key in product_name or product_name in key:
            return mod
    return None


def build_module_revenue(excel_path: str | Path) -> dict[str, dict[str, Any]]:
    """Aggregate YTD revenue (and optional growth) by module."""
    excel_path = Path(excel_path)
    data_rows, growth_mode = _read_export_sheet(excel_path)
    taxonomy = load_taxonomy()
    module_to_team = taxonomy.get("module_to_team", {})
    by_module: dict[str, dict[str, Any]] = defaultdict(
        lambda: {"revenue_eur": 0.0, "growth_eur": 0.0, "products": []}
    )
    unmapped: list[tuple[str, float]] = []

    for row in data_rows:
        if not row:
            continue
        product_full = str(row[0]).strip() if row[0] is not None else ""
        if not product_full or not product_full.startswith("A0"):
            continue

        product_name = extract_product_name(product_full)
        ytd = _to_float(row[-1])
        if ytd is None:
            continue
        if ytd <= 0 and product_full != "A01040 Consulting Customized":
            continue

        module = resolve_module(product_name)
        if not module:
            unmapped.append((product_name, ytd))
            continue

        growth = 0.0
        if growth_mode and len(row) > 8:
            g = _to_float(row[8])
            if g is not None:
                growth = g

        entry = by_module[module]
        entry["revenue_eur"] += ytd
        entry["growth_eur"] += growth
        entry["products"].append({"product": product_name, "ytd_eur": ytd})

    result: dict[str, dict[str, Any]] = {}
    for module, data in by_module.items():
        result[module] = {
            "module": module,
            "team": module_to_team.get(module, "Other"),
            "revenue_eur": round(data["revenue_eur"], 2),
            "growth_eur": round(data["growth_eur"], 2),
            "product_count": len(data["products"]),
        }

    return {"by_module": result, "unmapped_products": unmapped}
