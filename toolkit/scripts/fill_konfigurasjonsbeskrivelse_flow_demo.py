#!/usr/bin/env python3
"""Fill consultant Excel 'Konfigurasjonsbeskrivelse prosessmodell' from 360° Flow demo JSON (Norwegian)."""

from __future__ import annotations

import json
import shutil
from datetime import date
import sys
from pathlib import Path

import openpyxl

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))
from konfigurasjonsbeskrivelse_aktiviteter_no import write_aktiviteter_sheet

REPO = Path(__file__).resolve().parents[2]
PM_JSON = REPO / "projects/process-platform/flow/manuscripts/Demo setup/Behandle søknad m_klagerett.json"
CL_JSON = (
    REPO
    / "projects/process-platform/flow/manuscripts/Demo setup/Sjekkliste_for_søknadsbehandling_med_klagerett_fcba25b7-4d74-49ba-92c0-a03909c51d02.json"
)
TEMPLATE = Path(
    r"C:\Users\MortenJensen\Downloads\77392 Konfigurasjonsbeskrivelse prosessmodell 633483_4_0.xlsx"
)
OUTPUT = Path(
    r"C:\Users\MortenJensen\Downloads\77392 Konfigurasjonsbeskrivelse prosessmodell 633483_4_0 - 360 Flow demo NO.xlsx"
)

WU_TO_PROSESSDEF = {
    k.upper(): v
    for k, v in {
        "E1A5B6C7-D8F9-4012-9E3F-4B6C8D0A2E4F": "Final processing",
        "89e0ed8c-7391-4202-9f4d-d9d2d631c37a": (
            "Document Creation Process with Review, approval, dispatch and Milestone"
        ),
        "9E7A8B38-C249-4CD8-A95C-0EDD06ADB727": "MajorDeficiencyLetter",
        "10C6E5A9-C167-459B-B556-A0973A5E31A4": "Autosend Application Receipt to applicant",
        "6F8C2A1B-4E7D-4C9A-A3F1-2B5D8E9C0A1F": "create single document",
    }.items()
}

DOC_CAT_NAMES = {
    433: "Tilleggsinformasjon - mindre mangel",
    434: "Tilleggsinformasjon - vesentlig mangel",
    437: "Mangelbrev - mindre mangel",
    438: "Mangelbrev - vesentlig mangel",
    99002: "Vedtaksbrev",
    99024: "Klage",
    99012: "Søknad",
    111: "Dokument ut",
}

STAGE_TYPE_NO = "Lokal saksbehandlingsfase"
STAGE_CLASS_NO = "GEN - Generell"
DEADLINE_CALC_NO = "Beregn fra Start/Slutt milepælene"
CASE_TYPE_LABEL = "Sak"

# Consultant example milestone order (AutomationID prefix)
MILESTONE_AUTOMATION_ORDER = [
    "386b8fa1",
    "ea622bde",
    "5feeb7e2",
    "424ed2d6",
    "bc5007cc",
    "4e6f1342",
    "05b4d083",
    "bc78bdc4",
    "7c058a8b",
]
SKIP_MILESTONE_AUTOMATION = {"78be38e9"}  # duplicate tilleggsdokumentasjon mottatt

MILESTONE_LABEL_ALIASES = {
    "Søknad komplett": "Komplett søknad",
}

AUTOSEND_WU = "10C6E5A9-C167-459B-B556-A0973A5E31A4"

# From template sheet Kildegrunnlag (Norwegian customer-facing labels)
CONCURRENCY_NO = {
    1: "Bare én eksisterer",
    2: "Bare én eksisterer, status er ikke lukket",
    3: "Flere eksisterer, både åpne og lukkede",
}

MILESTONE_STATUS_NO = {
    1: "Nådd (UF)",
    2: "Ikke nådd (IUF)",
    3: "Planlagt (PL)",
    4: "Under behandling (B)",
    5: "Ikke relevant (NR)",
    6: "Fjernet (FJ)",
    7: "Flyttet (FL)",
}
MILESTONE_STATUS_FROM_DESC = {
    "Reached": "Nådd (UF)",
    "Not reached": "Ikke nådd (IUF)",
}

PROCESS_STATUS_NO = {
    1: "Planlagt",
    2: "Aktiv",
    3: "Avsluttet",
}
PROCESS_STATUS_FROM_DESC = {
    "Active": "Aktiv",
    "Closed": "Avsluttet",
    "Planned": "Planlagt",
}

TIME_CALC = {
    "0": "Ingen",
    "1": "Start tidsberegning",
    "2": "Sett tidsberegning på pause",
    "3": "Gjenoppta tidsberegning",
    "4": "Stopp tidsberegning",
}


def concurrency_label(value: object) -> str | None:
    if value is None:
        return None
    try:
        return CONCURRENCY_NO.get(int(value), str(value))
    except (TypeError, ValueError):
        return str(value)


def milestone_status_label(props: dict) -> str | None:
    recno = props.get("MilestoneStatusRecno")
    if recno is not None:
        label = MILESTONE_STATUS_NO.get(int(recno))
        if label:
            return label
    desc = props.get("MilestoneStatusDesc")
    if desc:
        return MILESTONE_STATUS_FROM_DESC.get(str(desc), desc)
    return None


def process_status_label(props: dict) -> str | None:
    recno = props.get("ProcessStatusRecno")
    if recno is not None:
        label = PROCESS_STATUS_NO.get(int(recno))
        if label:
            return label
    desc = props.get("ProcessStatusDescription")
    if desc:
        return PROCESS_STATUS_FROM_DESC.get(str(desc), desc)
    return None


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def milestone_type_label(desc: str | None) -> str | None:
    if not desc:
        return None
    clean = str(desc).lstrip("*").strip()
    return MILESTONE_LABEL_ALIASES.get(clean, clean)


def milestone_lookup(pm: dict) -> dict[int, str]:
    out: dict[int, str] = {}
    for stage in pm["ProcessModelData"]["StageConfigs"]:
        for ms in stage.get("StageMileStoneConfigs") or []:
            p = ms.get("Properties") or {}
            recno = p.get("MilestoneTypeRecno")
            desc = milestone_type_label(p.get("MilestoneTypeDesc") or p.get("Title"))
            if recno is not None and desc:
                out[int(recno)] = desc
    return out


def format_milestone_doc_categories(recnos: list[int]) -> str | None:
    if not recnos:
        return None
    if 433 in recnos and 434 in recnos:
        return (
            "Tilleggsinformasjon - mindre mangel, "
            "Tilleggsinformasjon - vesentlig mangel"
        )
    if 99002 in recnos:
        return DOC_CAT_NAMES[99002]
    names: list[str] = []
    for r in recnos:
        if r in (111, 60006) and 99002 in recnos:
            continue
        label = DOC_CAT_NAMES.get(r)
        if label and label not in names:
            names.append(label)
    return ", ".join(names) if names else None


def milestone_sort_key(ms: dict) -> tuple[int, str]:
    p = ms.get("Properties") or {}
    aid = (p.get("AutomationID") or "").lower()
    for idx, prefix in enumerate(MILESTONE_AUTOMATION_ORDER):
        if aid.startswith(prefix):
            return (idx, p.get("Title") or "")
    return (999, p.get("Title") or "")


def iter_milestones(pm: dict):
    for stage in pm["ProcessModelData"]["StageConfigs"]:
        stage_title = (stage.get("Properties") or {}).get("Title")
        items = list(stage.get("StageMileStoneConfigs") or [])
        items.sort(key=milestone_sort_key)
        for ms in items:
            p = ms.get("Properties") or {}
            aid = (p.get("AutomationID") or "").lower()
            if any(aid.startswith(s) for s in SKIP_MILESTONE_AUTOMATION):
                continue
            yield stage_title, ms


def process_sort_key(stage: dict, proc: dict) -> tuple:
    sp = stage.get("Properties") or {}
    p = proc.get("Properties") or {}
    return (sp.get("SortId") or 0, p.get("SortId") or 0)


def target_milestone_label(props: dict) -> str | None:
    return milestone_type_label(props.get("TargetMilestoneTypeDescription"))


def set_row(ws, row: int, values: dict[int, object]) -> None:
    for col, val in values.items():
        ws.cell(row=row, column=col, value=val)


def fill_forside(ws, pmc_title: str) -> None:
    ws.cell(row=3, column=2, value=f"Prosessmodell {pmc_title}")
    ws.cell(row=5, column=3, value="360° Flow standard demo (intern referanse)")
    ws.cell(row=6, column=3, value=date.today().isoformat())
    ws.cell(row=13, column=2, value="V0.1")
    ws.cell(row=13, column=3, value=date.today().isoformat())
    ws.cell(row=13, column=4, value="Morten Jensen / Cursor")
    ws.cell(row=13, column=5, value="Utfylt fra demo prosessmodell v9 og sjekkliste (NO JSON)")


def fill_prosessmodell(ws, pm: dict) -> None:
    props = pm["ProcessModelData"]["Properties"]
    title = props.get("Title") or pm.get("Title")
    row = 6
    set_row(
        ws,
        row,
        {
            2: title,
            3: None,
            4: title,
            5: CASE_TYPE_LABEL,
            8: "Nei",
        },
    )


def fill_faser(ws, pm: dict, checklist_name: str) -> None:
    row = 6
    pmc_title = pm["ProcessModelData"]["Properties"].get("Title") or pm.get("Title")
    for stage in pm["ProcessModelData"]["StageConfigs"]:
        sp = stage.get("Properties") or {}
        is_appeal = sp.get("Title") == "Klagebehandling"
        set_row(
            ws,
            row,
            {
                2: pmc_title,
                3: sp.get("Title"),
                4: STAGE_TYPE_NO,
                5: sp.get("Description"),
                6: sp.get("AdminNotes"),
                7: sp.get("CalculationDays"),
                8: sp.get("SortId"),
                9: sp.get("AutomationID"),
                10: concurrency_label(sp.get("Concurrency")),
                11: DEADLINE_CALC_NO,
                12: STAGE_CLASS_NO,
                13: None if is_appeal else checklist_name,
                14: "Ja" if sp.get("InitiateAsPlannedUpcoming") else "Nei",
                15: "Ja" if sp.get("AllowManualAdd") else "Nei",
                16: "Ja" if sp.get("AutoGenerate") else "Nei",
            },
        )
        row += 1


def fill_milepaler(ws, pm: dict) -> None:
    row = 6
    pmc_title = pm["ProcessModelData"]["Properties"].get("Title") or pm.get("Title")
    for stage_title, ms in iter_milestones(pm):
        p = ms.get("Properties") or {}
        cats = p.get("DocumentCategoryRecnos") or []
        cat_text = format_milestone_doc_categories(cats)
        set_row(
            ws,
            row,
            {
                2: pmc_title,
                3: stage_title,
                4: milestone_type_label(p.get("MilestoneTypeDesc")),
                5: p.get("Title"),
                6: p.get("Description"),
                7: p.get("AdminNotes"),
                8: p.get("CalculationDays"),
                9: p.get("SortId"),
                10: p.get("AutomationID"),
                11: milestone_status_label(p),
                12: concurrency_label(p.get("Concurrency")),
                13: cat_text,
                14: "Ja" if p.get("IsGlobal") else "Nei",
                15: "Ja" if p.get("AutoGenerate") else "Nei",
                16: "Ja" if p.get("ShouldRecalculate") else "Nei",
                17: "Ja" if p.get("IsStartMilestone") else "Nei",
                18: "Ja" if p.get("AllowManualAdd") else "Nei",
                19: TIME_CALC.get(
                    str(p.get("TimeCalculationBehaviorString")),
                    p.get("TimeCalculationBehaviorString"),
                ),
                20: None,
            },
        )
        row += 1


def fill_prosesser(ws, pm: dict) -> None:
    row = 6
    pmc_title = pm["ProcessModelData"]["Properties"].get("Title") or pm.get("Title")
    for stage in pm["ProcessModelData"]["StageConfigs"]:
        stage_title = (stage.get("Properties") or {}).get("Title")
        processes = sorted(
            stage.get("StageProcessConfigs") or [],
            key=lambda x: (x.get("Properties") or {}).get("SortId") or 0,
        )
        for proc in processes:
            p = proc.get("Properties") or {}
            wu = (p.get("WorkunitId") or "").upper()
            display = p.get("DisplayName")
            if stage_title == "Klagebehandling" and wu.startswith("6F8C2A1B"):
                display = "Mottatt klage"
            auto_run = retry = None
            if wu == AUTOSEND_WU.upper():
                auto_run = retry = "Ja"
            set_row(
                ws,
                row,
                {
                    2: pmc_title,
                    3: stage_title,
                    4: display,
                    5: p.get("ProcessTypeDesc"),
                    6: WU_TO_PROSESSDEF.get(wu, wu),
                    7: process_status_label(p),
                    8: p.get("Description"),
                    9: p.get("AdminNotes"),
                    10: concurrency_label(p.get("Concurrency")),
                    11: target_milestone_label(p),
                    12: p.get("SortId"),
                    13: p.get("AutomationID"),
                    14: "Ja" if p.get("ShowAsUpcomingProcess") else "Nei",
                    15: "Ja" if p.get("AllowManualAdd") else "Nei",
                    16: "Ja" if p.get("AutoGenerate") else "Nei",
                    17: "Ja" if p.get("ShowActivePassiveButton") else "Nei",
                    18: "Ja" if p.get("ShowClosePhaseButton") else "Nei",
                    19: auto_run,
                    20: retry,
                    21: "Nei" if not p.get("PreventProcessCancellation") else "Ja",
                    22: None,
                },
            )
            row += 1


CHECKPOINT_ANSWER_LABELS = {
    1: "Ja/Nei",
    2: "Fullført/Ikke fullført",
    3: "Ja/Nei/Ikke relevant",
}


def fill_sjekklister(ws, cl: dict, ms_lookup: dict[int, str]) -> None:
    row = 6
    name = cl.get("Name")
    set_row(
        ws,
        row,
        {
            2: name,
            3: "Alle" if cl.get("IsAllCaseType") == -1 else cl.get("ToCaseType"),
        },
    )
    row += 1
    cp_by_id = {cp.get("CheckpointId"): cp.get("Name") for cp in cl.get("CheckPoints") or []}
    for cp in cl.get("CheckPoints") or []:
        outcome = cp.get("Outcome")
        outcome_title = outcome_desc = outcome_type = None
        if outcome:
            ot = outcome.get("OutcomeText") or {}
            outcome_title = ot.get("Title")
            outcome_desc = ot.get("Description")
            outcome_type = outcome.get("OutcomeType")
        mt = cp.get("MilestoneType")
        mt_label = ms_lookup.get(int(mt), str(mt)) if mt else None
        parent_idx = cp.get("ParentIndex")
        parent_label = cp_by_id.get(parent_idx) if parent_idx else None
        answer_opt = cp.get("CheckpointAnswerOption")
        set_row(
            ws,
            row,
            {
                6: cp.get("CheckpointId"),
                7: cp.get("Name"),
                8: parent_label,
                9: cp.get("Description"),
                10: mt_label,
                12: cp.get("Theme"),
                13: CHECKPOINT_ANSWER_LABELS.get(answer_opt, answer_opt),
                14: outcome_desc,
                15: outcome_type,
                16: outcome_title,
                17: outcome_desc,
            },
        )
        row += 1


def fill_innstillinger(ws) -> None:
    """Demo-specific toggles (append after template examples)."""
    row = 47
    extras = [
        ("NVE", "På", "360° Flow / NVE-prosessbibliotek (demo)"),
        ("FFUI.CaseCreation", "Av", "Demo wiki: deaktiver for forventet saksopprettelse"),
    ]
    for toggle, state, comment in extras:
        set_row(ws, row, {2: toggle, 3: state, 4: comment})
        row += 1
    set_row(
        ws,
        26,
        {
            2: "Prosessrelaterte verdier",
            3: "Ja",
            4: "Demo: aktiver dokumentkategorier 433, 434, 437, 438, 99002, 99024; importer v9 prosessmodell og sjekkliste",
        },
    )


def fill_oversikt_fase(ws, title: str, stage: dict) -> None:
    ws.cell(row=2, column=2, value=title)
    row = 6
    processes = sorted(
        stage.get("StageProcessConfigs") or [],
        key=lambda x: (x.get("Properties") or {}).get("SortId") or 0,
    )
    for proc in processes:
        p = proc.get("Properties") or {}
        ws.cell(row=row, column=2, value=p.get("SortId"))
        row += 2  # consultant skeleton: sort id only, blank narrative columns


def main() -> None:
    if not TEMPLATE.is_file():
        raise SystemExit(f"Template not found: {TEMPLATE}")
    if not PM_JSON.is_file() or not CL_JSON.is_file():
        raise SystemExit("Demo JSON sources missing in repo")

    pm = load_json(PM_JSON)
    cl = load_json(CL_JSON)
    ms_lookup = milestone_lookup(pm)

    shutil.copy2(TEMPLATE, OUTPUT)
    wb = openpyxl.load_workbook(OUTPUT)

    fill_forside(wb["Forside"], pm.get("Title") or "")
    fill_prosessmodell(wb["Prosessmodell"], pm)
    fill_faser(wb["Saksbehandlingsfaser"], pm, cl.get("Name") or "")
    fill_milepaler(wb["Milep\u00e6ler"], pm)
    fill_prosesser(wb["Prosesser"], pm)
    fill_sjekklister(wb["Sjekklister"], cl, ms_lookup)
    fill_innstillinger(wb["Innst. og kodetabeller"])
    write_aktiviteter_sheet(
        wb["Aktiviteter"],
        pm,
        wu_to_prosessdef=WU_TO_PROSESSDEF,
        process_sort_key=process_sort_key,
    )

    stages = pm["ProcessModelData"]["StageConfigs"]
    if len(stages) >= 1:
        fill_oversikt_fase(
            wb["Oversikt Fase 1"],
            f"Oversikt fase 1 - {(stages[0].get('Properties') or {}).get('Title')}",
            stages[0],
        )
    if len(stages) >= 2:
        fill_oversikt_fase(
            wb["Oversikt Fase 2"],
            f"Oversikt fase 2 - {(stages[1].get('Properties') or {}).get('Title')}",
            stages[1],
        )

    wb.save(OUTPUT)
    repo_copy = REPO / "outputs" / "360_Flow_demo_NO_konfigurasjonsbeskrivelse.xlsx"
    repo_copy.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUTPUT, repo_copy)
    print(f"Wrote {OUTPUT}")
    print(f"  Processes: {sum(len(s.get('StageProcessConfigs') or []) for s in stages)}")
    print(f"  Milestones: {sum(len(s.get('StageMileStoneConfigs') or []) for s in stages)}")
    print(f"  Checkpoints: {len(cl.get('CheckPoints') or [])}")


if __name__ == "__main__":
    main()
