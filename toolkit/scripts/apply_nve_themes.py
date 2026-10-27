"""Apply agreed Tema values to NVE behov Excel sheets (column A, In progress)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

# Excel row number -> Tema (header row 1)
THEMES: dict[int, str] = {
    2: "Admin – prosessmodell og komponenter",
    3: "Kodetabeller og standarddata",
    4: "Webdeler og portefølje",
    5: "Sak og sakstype",
    6: "Kodetabeller og standarddata",
    7: "Admin – brukervennlighet og klarspråk",
    8: "Kodetabeller og standarddata",
    9: "Søk og filtrering",
    10: "Søk og filtrering",
    11: "Søk og filtrering",
    12: "Frister og tidsberegning",
    13: "Kodetabeller og standarddata",
    14: "Prosessvisning og saksflyt",
    15: "Admin – prosessmodell og komponenter",
    16: "Komponenter i prosess",
    17: "Komponenter i prosess",
    18: "Admin – prosessmodell og komponenter",
    19: "Vedtak, dispensasjon og saksfremlegg",
    20: "Frister og tidsberegning",
    21: "Søk og filtrering",
    22: "Søk og filtrering",
    23: "Prosessvisning og saksflyt",
    24: "Vedtak, dispensasjon og saksfremlegg",
    25: "Kodetabeller og standarddata",
    26: "Admin – brukervennlighet og klarspråk",
    27: "Webdeler og portefølje",
    28: "Komponenter i prosess",
    29: "Admin – prosessmodell og komponenter",
    30: "Admin – brukervennlighet og klarspråk",
    31: "Admin – prosessmodell og komponenter",
    32: "Prosessvisning og saksflyt",
    33: "Admin – brukervennlighet og klarspråk",
    34: "Admin – prosessmodell og komponenter",
    35: "Sak og sakstype",
    36: "Admin – prosessmodell og komponenter",
    37: "Admin – prosessmodell og komponenter",
    38: "Komponenter i prosess",
    39: "Admin – prosessmodell og komponenter",
    40: "Kodetabeller og standarddata",
    41: "Kodetabeller og standarddata",
    42: "Søk og filtrering",
    43: "Vedtak, dispensasjon og saksfremlegg",
    44: "Vedtak, dispensasjon og saksfremlegg",
    45: "Frister og tidsberegning",
    46: "Admin – prosessmodell og komponenter",
    47: "Dokument, filer og visning",
    48: "Ekspedering, e-post og Outlook",
    49: "Komponenter i prosess",
    50: "Dokument, filer og visning",
    51: "Frister og tidsberegning",
    52: "Admin – prosessmodell og komponenter",
    53: "Prosessvisning og saksflyt",
    54: "Vedtak, dispensasjon og saksfremlegg",
    55: "Komponenter i prosess",
    56: "Komponenter i prosess",
    57: "Kodetabeller og standarddata",
    58: "Kodetabeller og standarddata",
    59: "Admin – prosessmodell og komponenter",
    60: "Ekspedering, e-post og Outlook",
    61: "Ekspedering, e-post og Outlook",
    62: "Ekspedering, e-post og Outlook",
    63: "Admin – brukervennlighet og klarspråk",
    64: "Sjekkliste",
    65: "Webdeler og portefølje",
    66: "Admin – brukervennlighet og klarspråk",
    67: "Komponenter i prosess",
    68: "Admin – prosessmodell og komponenter",
    69: "Admin – prosessmodell og komponenter",
    70: "Frister og tidsberegning",
    71: "Admin – prosessmodell og komponenter",
    72: "Søk og filtrering",
    73: "Sjekkliste",
    74: "Webdeler og portefølje",
    75: "Komponenter i prosess",
    76: "Søk og filtrering",
    77: "Prosessvisning og saksflyt",
    78: "Integrasjon og produktparitet",
    79: "Integrasjon og produktparitet",
    80: "Sjekkliste",
    81: "Sjekkliste",
    82: "Integrasjon og produktparitet",
    83: "Sjekkliste",
    84: "Sjekkliste",
    85: "Sjekkliste",
    86: "Sjekkliste",
    87: "Sjekkliste",
    88: "Sjekkliste",
    89: "Sjekkliste",
    90: "Frister og tidsberegning",
    91: "Sjekkliste",
    92: "Sjekkliste",
    93: "Webdeler og portefølje",
    94: "Integrasjon og produktparitet",
    95: "Integrasjon og produktparitet",
    96: "Integrasjon og produktparitet",
    97: "Admin – prosessmodell og komponenter",
    98: "Sjekkliste",
    99: "Frister og tidsberegning",
    100: "Admin – prosessmodell og komponenter",
    101: "Admin – prosessmodell og komponenter",
    102: "Admin – prosessmodell og komponenter",
    103: "Admin – prosessmodell og komponenter",
    104: "Komponenter i prosess",
    105: "Frister og tidsberegning",
    106: "Frister og tidsberegning",
    107: "Søk og filtrering",
    108: "Dokument, filer og visning",
    109: "Sjekkliste",
    110: "Admin – prosessmodell og komponenter",
    111: "Ekspedering, e-post og Outlook",
    112: "Ekspedering, e-post og Outlook",
    113: "Komponenter i prosess",
    114: "Komponenter i prosess",
    115: "Admin – prosessmodell og komponenter",
}

OUTPUT = Path(r"C:\Users\MortenJensen\OneDrive - Tieto\NVE behov - DevOps status.xlsx")
SOURCE = Path(r"C:\Users\MortenJensen\OneDrive - Tieto\NVE behov.xlsx")


def apply(path: Path) -> None:
    wb = load_workbook(path)
    ws = wb["In progress"]
    for row_num, tema in THEMES.items():
        ws.cell(row=row_num, column=1, value=tema)
    wb.save(path)
    print(f"Updated Tema on {len(THEMES)} rows in {path}")


def main() -> None:
    for path in (OUTPUT, SOURCE):
        if path.is_file():
            apply(path)
        else:
            print(f"Skip missing: {path}")


if __name__ == "__main__":
    main()
