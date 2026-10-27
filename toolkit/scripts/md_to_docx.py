"""Convert a simple markdown file to Word (.docx). Supports headings, tables, bold, lists."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt


_LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def add_rich_text(paragraph, text: str) -> None:
    pos = 0
    for m in _LINK_RE.finditer(text):
        if m.start() > pos:
            _add_bold_runs(paragraph, text[pos : m.start()])
        label, url = m.group(1), m.group(2)
        run = paragraph.add_run(label)
        run.underline = True
        try:
            run.hyperlink.address = url
        except AttributeError:
            run = paragraph.add_run(f"{label} ({url})")
        pos = m.end()
    if pos < len(text):
        _add_bold_runs(paragraph, text[pos:])


def _add_bold_runs(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part:
            paragraph.add_run(part)


def is_table_row(line: str) -> bool:
    stripped = line.strip()
    return stripped.startswith("|") and stripped.endswith("|") and stripped.count("|") >= 2


def is_separator_row(line: str) -> bool:
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return all(re.fullmatch(r":?-{3,}:?", c or "---") for c in cells)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and is_table_row(lines[i]):
        if not is_separator_row(lines[i]):
            rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
        i += 1
    return rows, i


def convert_markdown_to_docx(source: Path, destination: Path) -> None:
    text = source.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.LINE)
            i += 1
            continue

        if stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=0)
            i += 1
            continue

        if stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=1)
            i += 1
            continue

        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
            i += 1
            continue

        if is_table_row(stripped):
            rows, i = parse_table(lines, i)
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Table Grid"
                for r_idx, row in enumerate(rows):
                    for c_idx, cell_text in enumerate(row):
                        cell = table.rows[r_idx].cells[c_idx]
                        cell.text = ""
                        p = cell.paragraphs[0]
                        add_rich_text(p, cell_text)
            continue

        if re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            add_rich_text(p, re.sub(r"^\d+\.\s+", "", stripped))
            i += 1
            continue

        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_rich_text(p, stripped[2:])
            i += 1
            continue

        p = doc.add_paragraph()
        add_rich_text(p, stripped)
        i += 1

    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert markdown to Word docx")
    parser.add_argument("source", type=Path)
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        help="Output .docx path (default: same name with .docx)",
    )
    args = parser.parse_args()
    output = args.output or args.source.with_suffix(".docx")
    convert_markdown_to_docx(args.source, output)
    print(output)


if __name__ == "__main__":
    main()
