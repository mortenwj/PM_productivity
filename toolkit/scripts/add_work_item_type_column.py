"""Insert Work item type column (after Beskrivelse) in NVE behov - DevOps status.xlsx."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

# Reuse parsers from build script
sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_nve_behov_split import ADO_CACHE, OUTPUT, extract_work_item_id, load_source  # noqa: E402

TYPE_COLUMN = "Work item type"


def devops_link_column_index(header: list) -> int:
    if "DevOps link" in header:
        return header.index("DevOps link")
    if "Status" in header:
        return header.index("Status")
    return 3


def work_item_type_from_link(link_cell: object, ado: dict) -> str:
    wid = extract_work_item_id(link_cell)
    if wid is None and link_cell:
        s = str(link_cell)
        m = re.match(r"^(User Story|Bug|Feature)\s+\d+", s, re.I)
        if m:
            return m.group(1).title().replace("User story", "User Story")
    if wid is None:
        return ""
    info = ado.get(str(wid)) or ado.get(wid)
    if info and info.get("type"):
        return info["type"]
    return ""


def apply(path: Path, ado: dict) -> None:
    wb = load_workbook(path)
    ws = wb["In progress"]
    header = [c.value for c in ws[1]]

    if TYPE_COLUMN in header:
        type_col = header.index(TYPE_COLUMN) + 1
    else:
        ws.insert_cols(3)
        ws.cell(row=1, column=3, value=TYPE_COLUMN)
        type_col = 3
        header = [c.value for c in ws[1]]

    link_col = devops_link_column_index(header) + 1

    for row_num in range(2, ws.max_row + 1):
        link_val = ws.cell(row=row_num, column=link_col).value
        ws.cell(row=row_num, column=type_col, value=work_item_type_from_link(link_val, ado))

    wb.save(path)
    print(f"Updated {path} ({ws.max_row - 1} data rows)")


def main() -> int:
    if not ADO_CACHE.is_file():
        print(f"Missing cache: {ADO_CACHE}", file=sys.stderr)
        return 1
    ado = json.loads(ADO_CACHE.read_text(encoding="utf-8"))
    if not OUTPUT.is_file():
        print(f"Missing output: {OUTPUT}", file=sys.stderr)
        return 1
    path = Path(load_source(OUTPUT))
    apply(path, ado)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
