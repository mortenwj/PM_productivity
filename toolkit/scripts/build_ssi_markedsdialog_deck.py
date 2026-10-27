#!/usr/bin/env python3
"""Build grundstruktur for SSI markedsdialog deck (dansk)."""

from __future__ import annotations

import shutil
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt

DECK_PATH = Path(r"C:\Users\MortenJensen\OneDrive - Tieto\SSI markedsdialog.pptx")

# Layout indices from corporate template in this file
LAYOUT_COVER = 1
LAYOUT_CONTENT = 16  # Content 3 — title + content
LAYOUT_TWO_COL = 21
LAYOUT_CHAPTER = 10
LAYOUT_AGENDA = 5
LAYOUT_END = 44


def delete_slide(prs: Presentation, index: int) -> None:
    slide_id = prs.slides._sldIdLst[index]
    r_id = slide_id.rId
    prs.part.drop_rel(r_id)
    del prs.slides._sldIdLst[index]


def clear_slides(prs: Presentation) -> None:
    while len(prs.slides) > 0:
        delete_slide(prs, len(prs.slides) - 1)


def find_title_shape(slide):
    if slide.shapes.title:
        return slide.shapes.title
    for shape in slide.shapes:
        if shape.has_text_frame and "Title" in shape.name:
            return shape
    return None


def find_body_shapes(slide) -> list:
    bodies = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        name = shape.name
        if "Title" in name:
            continue
        if any(x in name for x in ("Footer", "Slide Number", "Date", "Picture")):
            continue
        if "Placeholder" in name or "Text Placeholder" in name or "Content Placeholder" in name:
            bodies.append(shape)
    return bodies


def set_title(slide, text: str) -> None:
    shape = find_title_shape(slide)
    if shape is None:
        raise RuntimeError("No title shape on slide")
    shape.text_frame.text = text


def set_bullets(shape, lines: list[str], font_size: int = 18) -> None:
    tf = shape.text_frame
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        if p.runs:
            p.runs[0].font.size = Pt(font_size)


def add_content_slide(prs, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_CONTENT])
    set_title(slide, title)
    bodies = find_body_shapes(slide)
    if not bodies:
        raise RuntimeError(f"No body on content slide: {title}")
    set_bullets(bodies[0], bullets)


def add_chapter_slide(prs, title: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_CHAPTER])
    set_title(slide, title)


def main() -> None:
    if not DECK_PATH.is_file():
        raise SystemExit(f"Deck not found: {DECK_PATH}")

    backup = DECK_PATH.with_name(DECK_PATH.stem + ".backup.pptx")
    shutil.copy2(DECK_PATH, backup)

    prs = Presentation(str(DECK_PATH))
    clear_slides(prs)

    today = date.today().strftime("%d. %B %Y").replace("September", "september")

    # 1 — Forside
    cover = prs.slides.add_slide(prs.slide_layouts[LAYOUT_COVER])
    set_title(cover, "Markedsdialog — procesunderstøttelse i Public 360°")
    bodies = find_body_shapes(cover)
    if len(bodies) >= 1:
        bodies[0].text_frame.text = "Statens Serum Institut · DBIS"
    if len(bodies) >= 2:
        bodies[1].text_frame.text = f"Case: rettighedsanmodning (PKU / destruktion, barn under 15 år)\n{today}"

    # 2
    add_content_slide(
        prs,
        "Jeres case — hvad vi har hørt",
        [
            "Forælder anmoder på vegne af barn under 15 år: indsigt i PKU-kort og destruktion",
            "DBIS — rettighedsanmodninger med identitet, forældremyndighed og intern høring",
            "I dag: ESDH + manuelt overblik (fx Borgerarket) — ikke krav om 1:1 genskabelse",
            "Målet: se hvordan processen bedst understøttes med færre manuelle registreringer",
        ],
    )

    # 3
    add_content_slide(
        prs,
        "SSI’s ønsker til demonstrationen",
        [
            "Modtagelse, sagsoprettelse og genbrug af oplysninger på sagen",
            "Regler og grene (identitet, samtykke, intern høring)",
            "Frister — beregning, pause ved afventning, påmindelser/rykkere",
            "Opgaver, ansvar, status og perioder med afventning (borger / organisation)",
            "Dokumenter: udkast, godkendelse, udsendelse, journalisering",
            "Afslutning, genåbning og overblik for sagsbehandler og ledelse",
            "Konfiguration: hvad kan SSI selv ændre uden leverandør?",
        ],
    )

    # 4
    add_content_slide(
        prs,
        "Public 360° + 360° Flow",
        [
            "Public 360°: sag, dokumenter, journalisering, roller og sikkerhed",
            "360° Flow: procesmodel på sagen — faser, milepæle, processer, tjeklister",
            "ProcessView: sagsbehandlerens overblik over hvor sagen er i forløbet",
            "Ikke et separat BPM-værktøj — styret progression inde i ESDH-kernen",
            "[Indsæt diagram: sag + Flow-lag]",
        ],
    )

    # 5 — To kolonner
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_TWO_COL])
    set_title(slide, "Fra manuelt overblik til styret forløb")
    cols = find_body_shapes(slide)
    if len(cols) >= 2:
        set_bullets(
            cols[0],
            [
                "I dag (typisk)",
                "Status og frister flere steder",
                "Manuel koordinering (mail, Excel)",
                "Svært overblik for ledelse",
            ],
        )
        set_bullets(
            cols[1],
            [
                "Med 360° Flow",
                "Én sag — status og milepæle samlet",
                "Regler, afventning og rykkere i modellen",
                "Arbejdsliste og portefølje efter fase/frist",
            ],
        )

    # 6
    add_content_slide(
        prs,
        "Case-mapping (forenklet)",
        [
            "1. Modtagelse og sagsoprettelse",
            "2. Identifikation og forældremyndighed",
            "3. Evt. samtykke ved fælles forældremyndighed → afvent borger",
            "4. Intern høring / destruktion (fx MS/PKU) → afvent forretning",
            "5. Opfølgning og rykker",
            "6. Afgørelse — udkast, godkendelse, udsendelse",
            "7. Afslutning — evt. genåbning ved destruktionsbekræftelse",
        ],
    )

    # 7
    add_content_slide(
        prs,
        "Integrationer og afgrænsninger",
        [
            "CPR og Digital Post (NgDP): platform/integration — dokumentation på sagen",
            "SSI-interne systemer (fx LIMS): integration eller arbejdsenhed; handling tilbage på sag",
            "Borgervendt digital blanket kan ligge uden for ESDH — sagen er stadig ankert",
            "Fuld automatisering af alle afdelingshøringer: projekt — i demo viser vi mønsteret",
        ],
    )

    # 8 — Demo-kapitel
    add_chapter_slide(prs, "Demo — sagsbehandling (ca. 60 min)")

    # 9 — Demo-agenda
    agenda = prs.slides.add_slide(prs.slide_layouts[LAYOUT_AGENDA])
    set_title(agenda, "Demo-agenda")
    bodies = find_body_shapes(agenda)
    if bodies:
        set_bullets(
            bodies[0],
            [
                "Forventningsafstemning — mønstre, ikke fuld 22-trins genskabelse",
                "Opret/åbn sag — journaliser henvendelse (rettighedsanmodning)",
                "ProcessView — faser, milepæle, parallelle processer",
                "Modtagelse + tjekliste (identitet, parter)",
                "Gren: manglende oplysninger / samtykke — afvent borger, frist pause",
                "Intern afventning + påmindelse (analog til MS/PKU-høring)",
                "Afgørelse — skabelon, godkendelse, udsendelse",
                "Afslutning — kort om genåbning ved destruktionsbekræftelse",
            ],
            font_size=16,
        )

    # 10
    add_chapter_slide(prs, "Konfiguration (ca. 10 min)")

    # 11
    add_content_slide(
        prs,
        "Hvad kan SSI selv konfigurere?",
        [
            "Procesmodel — faser, processer, milepæle (import/eksport afhænger af miljø)",
            "Tjeklister knyttet til milepæle og kvalitetssikring",
            "Dokumentskabeloner og kategorier (afgørelse, henvendelser)",
            "Roller: hvem publicerer model vs. udfører processer på sagen",
            "Spørgsmål til jer: Hvad vil I typisk justere selv?",
        ],
    )

    # 12
    add_content_slide(
        prs,
        "Implementering og 360° Flow (pilot)",
        [
            "360° Flow (P00285) — kontrolleret pilot; demo viser modenhed og retning",
            "Opstart med standard byggesten og implementeringspakke (typisk op til 2 procesmodeller)",
            "Vækst: flere processer og dybere governance på samme platform",
            "Grundlæggende sporbarhed er ikke paywall — kapacitet og services følger kommerciel model",
        ],
    )

    # 13
    add_content_slide(
        prs,
        "Ærlige afgrænsninger i dagens møde",
        [
            "Demo bruger referenceproces — fortælling følger jeres PKU/rettighedscase",
            "Ikke alle Borgerark-kolonner vises — behov dækkes via status/tjekliste/processer",
            "SSI-specifik LIMS/PKU-integration kræver afklaring i implementeringsprojekt",
            "Detaljer i pilot vs. generel tilgængelighed afstemmes efter markedsdialogen",
        ],
    )

    # 14 — Afslutning
    end = prs.slides.add_slide(prs.slide_layouts[LAYOUT_END])
    set_title(end, "Spørgsmål og næste skridt")
    bodies = find_body_shapes(end)
    if bodies:
        bodies[0].text_frame.text = "Tak for dialogen.\n\n[Kontaktpersoner · opfølgning markedsdialog]"

    prs.save(str(DECK_PATH))
    print(f"Saved {len(prs.slides)} slides to {DECK_PATH}")
    print(f"Backup: {backup}")


if __name__ == "__main__":
    main()
