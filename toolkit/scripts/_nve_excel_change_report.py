"""Report Excel-relevant changes from DevOps refresh (session snapshot)."""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

EXCEL = Path(r"C:\Users\MortenJensen\Downloads\NVE behov - DevOps status.xlsx")
CACHE = Path(__file__).resolve().parent / "nve_behov_ado_cache.json"

# DevOps state before refresh (from cache at session start, before live fetch)
OLD_STATE: dict[int, str] = {
    424420: "Closed",
    424871: "Resolved",
    425033: "Closed",
    425053: "New",
    425064: "Closed",
    425447: "Closed",
    425608: "Closed",
    426229: "New",
    426254: "New",
    426270: "Closed",
    426273: "Closed",
    426552: "Closed",
    427162: "Resolved",
    427213: "Resolved",
    427315: "Active",
    427319: "Rejected",
    427330: "Active",
    427331: "Rejected",
    427333: "Active",
    427338: "New",
    427339: "Closed",
    427366: "New",
    427613: "New",
    428100: "New",
    428128: "New",
    428130: "Active",
    428169: "New",
    428170: "New",
    428172: "New",
    428175: "New",
    428176: "New",
    428358: "New",
    428533: "Closed",
    428567: "Resolved",
    428570: "New",
    428605: "Resolved",
    428616: "Closed",
    428975: "Rejected",
    429044: "Active",
    429929: "Closed",
    429931: "Closed",
    429987: "Closed",
    429992: "New",
    429998: "New",
    430000: "New",
    430717: "Closed",
    431468: "New",
    431496: "Closed",
    431560: "Closed",
    431579: "Closed",
    431622: "Resolved",
    431676: "Rejected",
    431870: "New",
    433856: "Closed",
    435077: "New",
    435082: "New",
    435085: "New",
    435091: "New",
    435095: "New",
    435097: "New",
    435098: "New",
    435100: "New",
    435467: "New",
    435714: "New",
    435766: "New",
    436878: "New",
    436880: "New",
    437874: "New",
    437883: "New",
    438238: "New",
    438447: "Closed",
    438566: "Closed",
    439056: "Closed",
    439081: "New",
    439082: "New",
    439325: "New",
    439334: "New",
    439337: "New",
    439338: "New",
    439339: "New",
    439341: "New",
    439448: "New",
    440160: "New",
    440179: "New",
    440186: "New",
    440603: "New",
    440605: "New",
    440607: "New",
    440610: "New",
    440695: "Closed",
    441552: "New",
    441574: "New",
    441592: "New",
    441765: "New",
    441929: "Active",
}

NEW_IDS = [
    439092,
    442321,
    442324,
    442337,
    444336,
    444351,
    444357,
    444359,
    444360,
    444852,
    444916,
    444940,
    445534,
    445535,
    445912,
    445943,
    446237,
    446260,
    446510,
]


def wid_from(value: object) -> int | None:
    if value is None:
        return None
    m = re.search(r"(\d{5,7})", str(value))
    return int(m.group(1)) if m else None


def main() -> None:
    cache = json.loads(CACHE.read_text(encoding="utf-8"))
    wb = load_workbook(EXCEL, read_only=True, data_only=True)

    status_changes: list[tuple] = []
    for wid, old in OLD_STATE.items():
        new = cache.get(str(wid), {}).get("state", "")
        if old != new:
            title = cache.get(str(wid), {}).get("title", "")[:70]
            status_changes.append((wid, title, old, new))

    # Where each changed ID appears in Excel (ADO Status column before was `old`)
    sheet_rows: dict[int, list[str]] = {}
    for sn in ("Alle", "Under arbeid", "Done_Resolved", "Rejected"):
        if sn not in wb.sheetnames:
            continue
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        h = list(rows[0])
        if "DevOps link" not in h:
            continue
        li, si = h.index("DevOps link"), h.index("ADO Status") if "ADO Status" in h else None
        for rn, r in enumerate(rows[1:], start=2):
            r = list(r) + [None] * (len(h) - len(r))
            w = wid_from(r[li])
            if w:
                sheet_rows.setdefault(w, []).append(f"{sn} (row {rn})")

    new_in_excel: list[tuple] = []
    for wid in NEW_IDS:
        info = cache.get(str(wid), {})
        locs = sheet_rows.get(wid, [])
        if locs:
            new_in_excel.append((wid, info.get("state", ""), info.get("title", "")[:70], locs))

    wb.close()

    print("=== ADO Status changes (unique work items) ===")
    for wid, title, old, new in sorted(status_changes):
        loc = ", ".join(sheet_rows.get(wid, []))
        print(f"#{wid}  {old} -> {new}")
        print(f"    {title}")
        if loc:
            print(f"    Sheets: {loc}")
        print()

    print("=== New DevOps links populated in Excel (not in earlier cache) ===")
    for wid, st, title, locs in sorted(new_in_excel):
        print(f"#{wid}  {st}  —  {title}")
        print(f"    Sheets: {', '.join(locs)}")
        print()

    print("=== Summary ===")
    print(f"Status changes: {len(status_changes)} work items")
    print(f"New linked items in workbook: {len(new_in_excel)} of {len(NEW_IDS)} new cache IDs")


if __name__ == "__main__":
    main()
