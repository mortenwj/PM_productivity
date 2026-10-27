"""Diff current Excel vs reconstructed pre-refresh values (from session file compare)."""
from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path

from openpyxl import load_workbook

EXCEL = Path(r"C:\Users\MortenJensen\Downloads\NVE behov - DevOps status.xlsx")

# From old Downloads vs first refresh output (direct file compare)
STATUS_REVERT: dict[int, str | None] = {
    440160: "New",
    428567: "Resolved",
    431622: "Resolved",
    445912: None,
    446510: None,
}

SHEETS = ("Alle", "Under arbeid", "Done_Resolved", "Rejected")
COLS = ("Work item type", "DevOps link", "ADO Status", "Sprint", "Tags", "Opprettet", "Sist oppdatert")


def wid_from(v: object) -> int | None:
    if v is None:
        return None
    m = re.search(r"(\d{5,7})", str(v))
    return int(m.group(1)) if m else None


def main() -> None:
    wb = load_workbook(EXCEL, data_only=False)
    before = deepcopy(wb)

    for sn in SHEETS:
        if sn not in wb.sheetnames:
            continue
        ws_b = before[sn]
        rows = list(ws_b.iter_rows(min_row=1, max_row=ws_b.max_row))
        header = [c.value for c in rows[0]]
        if "DevOps link" not in header or "ADO Status" not in header:
            continue
        li = header.index("DevOps link")
        si = header.index("ADO Status")
        for row in rows[1:]:
            w = wid_from(row[li].value)
            if w in STATUS_REVERT:
                row[si].value = STATUS_REVERT[w]

    changes: list[tuple] = []
    for sn in SHEETS:
        if sn not in wb.sheetnames:
            continue
        ws_a, ws_b = wb[sn], before[sn]
        for r in range(2, ws_a.max_row + 1):
            for col_name in COLS:
                try:
                    ci = [ws_a.cell(1, c).value for c in range(1, ws_a.max_column + 1)].index(col_name) + 1
                except ValueError:
                    continue
                after = ws_a.cell(r, ci).value
                before_v = ws_b.cell(r, ci).value
                if str(after or "").strip() == str(before_v or "").strip():
                    continue
                link_ci = [
                    ws_a.cell(1, c).value for c in range(1, ws_a.max_column + 1)
                ].index("DevOps link") + 1
                w = wid_from(ws_a.cell(r, link_ci).value)
                changes.append((sn, r, w, col_name, before_v, after))

    wb.close()
    before.close()

    seen: set[tuple] = set()
    print("=== Changes in Excel from DevOps refresh (reconstructed before snapshot) ===\n")
    for sn, r, w, col, b, a in sorted(changes, key=lambda x: (x[2] or 0, x[3])):
        key = (w, col, str(b), str(a))
        if key in seen:
            continue
        seen.add(key)
        print(f"#{w or '?'} | {col}")
        print(f"  Before: {b!r}")
        print(f"  After:  {a!r}")
        print(f"  Example row: {sn} row {r}")
        print()

    print(f"Distinct value changes: {len(seen)}")


if __name__ == "__main__":
    main()
