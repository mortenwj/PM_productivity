"""Refresh ADO Status / Sprint / Tags (and optional dates) in NVE behov Excel in-place.

Preserves existing cell formatting; only updates values (and ensures DevOps hyperlinks).
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from build_nve_behov_split import (  # noqa: E402
    ADO_BASE,
    extract_work_item_id,
    format_sprint,
    format_tags,
    link_label,
    work_item_type,
)

DEFAULT_INPUT = Path(r"C:\Users\MortenJensen\Downloads\NVE behov - DevOps status.xlsx")
ADO_CACHE = SCRIPT_DIR / "nve_behov_ado_cache.json"

LINK_FONT = Font(color="0563C1", underline="single")

SHEETS_WITH_DEVOPS = ("Alle", "Under arbeid", "Done_Resolved", "Rejected")


def load_ado_map(cache_path: Path) -> dict[str, dict]:
    raw = json.loads(cache_path.read_text(encoding="utf-8"))
    out: dict[str, dict] = {}
    for k, v in raw.items():
        out[str(k)] = v
    return out


def parse_iso_date(value: object) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    s = str(value).strip()
    if not s:
        return None
    s = s.replace("Z", "+00:00")
    try:
        return datetime.fromisoformat(s)
    except ValueError:
        return None


def format_date_cell(dt: datetime | None) -> str | None:
    if dt is None:
        return None
    return dt.strftime("%Y-%m-%d")


def wid_from_cell(cell) -> int | None:
    wid = extract_work_item_id(cell.value)
    if wid is not None:
        return wid
    if cell.hyperlink and cell.hyperlink.target:
        m = re.search(r"/edit/(\d+)", cell.hyperlink.target, re.I)
        if m:
            return int(m.group(1))
    return None


def header_map(ws) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=1, column=col).value
        if val is not None and str(val).strip():
            mapping[str(val).strip()] = col
    return mapping


def update_sheet(ws, ado: dict[str, dict], stats: dict[str, int]) -> None:
    cols = header_map(ws)
    if "DevOps link" not in cols:
        return

    col_link = cols["DevOps link"]
    col_type = cols.get("Work item type")
    col_status = cols.get("ADO Status")
    col_sprint = cols.get("Sprint")
    col_tags = cols.get("Tags")
    col_created = cols.get("Opprettet")
    col_changed = cols.get("Sist oppdatert")

    for row in range(2, ws.max_row + 1):
        link_cell = ws.cell(row=row, column=col_link)
        wid = wid_from_cell(link_cell)
        if wid is None:
            continue

        info = ado.get(str(wid))
        if not info:
            stats["missing"] += 1
            if col_status:
                ws.cell(row=row, column=col_status).value = "(not found in ADO)"
            continue

        stats["updated"] += 1
        label = link_label(wid, ado)
        url = ADO_BASE.format(id=wid)

        link_cell.value = label
        if link_cell.hyperlink is None or not link_cell.hyperlink.target:
            link_cell.hyperlink = url
            link_cell.font = copy(LINK_FONT)
        elif link_cell.hyperlink.target != url:
            link_cell.hyperlink = url

        if col_type:
            ws.cell(row=row, column=col_type).value = work_item_type(info, wid, link_cell.value)
        if col_status:
            ws.cell(row=row, column=col_status).value = info.get("state", "")
        if col_sprint:
            ws.cell(row=row, column=col_sprint).value = format_sprint(info)
        if col_tags:
            ws.cell(row=row, column=col_tags).value = format_tags(info)

        created = parse_iso_date(info.get("created"))
        changed = parse_iso_date(info.get("changed"))
        if col_created and created:
            ws.cell(row=row, column=col_created).value = format_date_cell(created)
        if col_changed and changed:
            ws.cell(row=row, column=col_changed).value = format_date_cell(changed)


def enrich_cache_from_batch(cache_path: Path, batch_path: Path) -> None:
    from merge_ado_batch_into_cache import load_batch_array

    items = load_batch_array(batch_path)
    ado = json.loads(cache_path.read_text(encoding="utf-8"))
    for item in items:
        fields = item.get("fields") or {}
        wid = str(fields.get("System.Id") or item.get("id"))
        entry = ado.setdefault(wid, {})
        entry["type"] = fields.get("System.WorkItemType", entry.get("type", ""))
        entry["state"] = fields.get("System.State", entry.get("state", ""))
        entry["title"] = fields.get("System.Title", entry.get("title", ""))
        entry["iteration_path"] = fields.get("System.IterationPath") or ""
        entry["tags"] = fields.get("System.Tags") or ""
        entry["created"] = fields.get("System.CreatedDate") or ""
        entry["changed"] = fields.get("System.ChangedDate") or ""
    cache_path.write_text(json.dumps(ado, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="Update NVE DevOps status Excel in-place")
    parser.add_argument(
        "input",
        nargs="?",
        type=Path,
        default=DEFAULT_INPUT,
        help=f"Path to .xlsx (default: {DEFAULT_INPUT})",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        help="Save path (default: overwrite input)",
    )
    parser.add_argument(
        "--batch",
        type=Path,
        help="MCP/REST get_batch JSON file to refresh cache before update",
    )
    parser.add_argument("--cache", type=Path, default=ADO_CACHE)
    args = parser.parse_args()

    if not args.input.is_file():
        print(f"Input not found: {args.input}", file=sys.stderr)
        return 1

    if args.batch:
        enrich_cache_from_batch(args.cache, args.batch)

    if not args.cache.is_file():
        print(f"ADO cache missing: {args.cache}", file=sys.stderr)
        return 1

    ado = load_ado_map(args.cache)
    out_path = args.output or args.input

    wb = load_workbook(args.input)
    stats = {"updated": 0, "missing": 0}
    for name in SHEETS_WITH_DEVOPS:
        if name in wb.sheetnames:
            update_sheet(wb[name], ado, stats)
    wb.save(out_path)
    print(f"Wrote {out_path}")
    print(f"Rows refreshed: {stats['updated']}; missing in ADO: {stats['missing']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
