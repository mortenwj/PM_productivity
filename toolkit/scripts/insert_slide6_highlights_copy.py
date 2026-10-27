"""Duplicate slide 6, insert after slide 1, apply 6.12 wiki-aligned highlight text."""
from __future__ import annotations

import shutil
import sys
from copy import deepcopy
from datetime import datetime
from pathlib import Path

from pptx import Presentation

PPTX = Path(
    r"C:\Users\MortenJensen\OneDrive - Tieto\Shortcuts\SI Products 360PM - release presentations"
    r"\Internal release presentation public 360 Sept_Okt (6.12).pptx"
)

PILLAR_LABELS = {
    "Proactive decision support",
    "Expert domain knowledge",
    "Automated & transparent",
    "Intuitive design",
    "Sustainable & ethical",
}

LEFT_TEXT = [
    "360° Flow — Generally Available (6.12+)",
    "360° eSigning for Denmark",
    "360° Sign Sweden — reminders & org info",
    "360° eManager for Android",
    "Word Online standard texts (GA)",
    "FIKS IO — module separation",
    "Teams — unified Archive to 360°",
    "GEMI — email filters",
    "360° IFAR corrections (Norway)",
    "AFV + VeraPDF PDF/A conformity",
    "360° Search Assistant improvements",
]

RIGHT_TEXT = [
    "New Outlook in standard offering · Improved global search",
    "Case priority (all types) · Dependency queue monitoring · Tasks webpart defaults",
    "VeraPDF on upload · Full-Flex for Projects · Mimir source priority · eArchive Fiks IO",
]

TITLE_TEXT = "Highlights in this release — 6.12 (October 2026)"


def duplicate_slide(prs: Presentation, index: int):
    source = prs.slides[index]
    dest = prs.slides.add_slide(source.slide_layout)
    for shape in source.shapes:
        newel = deepcopy(shape.element)
        dest.shapes._spTree.insert_element_before(newel, "p:extLst")
    return dest


def move_slide(prs: Presentation, old_index: int, new_index: int) -> None:
    sld_id_lst = prs.slides._sldIdLst
    sld_id = sld_id_lst[old_index]
    sld_id_lst.remove(sld_id)
    sld_id_lst.insert(new_index, sld_id)


def set_shape_text(shape, text: str) -> None:
    if not shape.has_text_frame:
        raise ValueError(f"Shape {shape.name!r} has no text frame")
    shape.text_frame.paragraphs[0].text = text


def highlight_pentagons(slide):
    """Return highlight pentagons in slide z-order (matches template layout)."""
    shapes = []
    for sh in slide.shapes:
        if sh.name != "Pentagon 19" or not sh.has_text_frame:
            continue
        label = sh.text_frame.text.strip().replace("\u200b", "")
        if label in PILLAR_LABELS:
            continue
        shapes.append(sh)
    if len(shapes) != len(LEFT_TEXT) + len(RIGHT_TEXT):
        raise ValueError(f"Expected {len(LEFT_TEXT) + len(RIGHT_TEXT)} highlight shapes, got {len(shapes)}")
    return shapes[: len(LEFT_TEXT)], shapes[len(LEFT_TEXT) :]


def title_shape(slide):
    for sh in slide.shapes:
        if sh.name == "Title 1" and sh.has_text_frame and "Highlights" in sh.text_frame.text:
            return sh
    for sh in slide.shapes:
        if sh.has_text_frame and "Highlights in this release" in sh.text_frame.text:
            return sh
    raise ValueError("Title shape not found on highlights slide")


def main() -> int:
    if not PPTX.is_file():
        print(f"File not found: {PPTX}", file=sys.stderr)
        return 1

    backup = PPTX.with_name(
        PPTX.stem + f".backup-{datetime.now().strftime('%Y%m%d-%H%M%S')}" + PPTX.suffix
    )
    shutil.copy2(PPTX, backup)
    print(f"Backup: {backup}")

    prs = Presentation(str(PPTX))
    source_index = 5  # slide 6
    duplicate_slide(prs, source_index)
    new_index = len(prs.slides) - 1
    move_slide(prs, new_index, 1)  # after title slide -> slide 2

    slide = prs.slides[1]
    set_shape_text(title_shape(slide), TITLE_TEXT)
    left_shapes, right_shapes = highlight_pentagons(slide)
    for shape, text in zip(left_shapes, LEFT_TEXT):
        set_shape_text(shape, text)
    for shape, text in zip(right_shapes, RIGHT_TEXT):
        set_shape_text(shape, text)
    for sh in slide.shapes:
        if sh.name == "Slide Number Placeholder 3" and sh.has_text_frame:
            sh.text_frame.paragraphs[0].text = "2"
            break

    prs.save(str(PPTX))
    print(f"Updated: {PPTX}")
    print("Inserted wiki-aligned highlights copy as slide 2; original slide 6 unchanged (now slide 7).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
