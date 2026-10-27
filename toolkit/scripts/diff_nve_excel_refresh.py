"""Compare current NVE Excel to values before DevOps refresh (inferred from pre-update file).

Uses the first ADO batch export from the refresh session as 'after' truth and
reconstructs 'before' by reversing known edits captured when Downloads was
compared to the first updated copy, plus empty DevOps cells filled in.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

BATCH_AFTER = Path(
    r"C:\Users\MortenJensen\.cursor\projects\c-Users-MortenJensen-cursor-projects-PM-productivity\agent-tools\d8400dc6-3ac1-425e-977d-73f91cd69166.txt"
)
EXCEL = Path(r"C:\Users\MortenJensen\Downloads\NVE behov - DevOps status.xlsx")

# ADO Status in file before refresh (from direct old-vs-new file diff in session)
STATUS_BEFORE: dict[int, str | None] = {
    440160: "New",
    428567: "Resolved",
    431622: "Resolved",
    445912: None,
    446510: None,
}

SHEETS = ("Alle", "Under arbeid", "Done_Resolved", "Rejected")
TRACK_COLS = (
    "Work item type",
    "DevOps link",
    "ADO Status",
    "Sprint",
    "Tags",
    "Opprettet",
    "Sist oppdatert",
)


def load_batch(path: Path) -> dict[str, dict]:
    text = path.read_text(encoding="utf-8")
    start, end = text.find("["), text.rfind("]") + 1
    items = json.loads(text[start:end])

    def norm_sprint(path_val: str | None) -> str:
        if not path_val or path_val.strip() in ("360", ""):
            return "Ingen sprint"
        p = path_val.replace("/", "\\")
        if p == r"360\core\Process-platform":
            return "Ingen sprint"
        return p

    def norm_tags(t: str | None) -> str:
        if not t:
            return ""
        return "; ".join(x.strip() for x in str(t).split(";") if x.strip())

    out: dict[str, dict] = {}
    for it in items:
        f = it.get("fields") or {}
        wid = str(f.get("System.Id") or it.get("id"))
        wtype = f.get("System.WorkItemType", "")
        title = f.get("System.Title", "")
        out[wid] = {
            "type": wtype,
            "state": f.get("System.State", ""),
            "title": title,
            "sprint": norm_sprint(f.get("System.IterationPath")),
            "tags": norm_tags(f.get("System.Tags")),
            "link": f"{wtype} {wid}: {title}",
            "created": (f.get("System.CreatedDate") or "")[:10],
            "changed": (f.get("System.ChangedDate") or "")[:10],
        }
    return out


def wid_from(value: object) -> str | None:
    if value is None:
        return None
    m = re.search(r"(\d{5,7})", str(value))
    return m.group(1) if m else None


def main() -> int:
    excel_path = Path(sys.argv[1]) if len(sys.argv) > 1 else EXCEL
    batch_path = Path(sys.argv[2]) if len(sys.argv) > 2 else BATCH_AFTER

    if not excel_path.is_file():
        print(f"Missing Excel: {excel_path}", file=sys.stderr)
        return 1
    if not batch_path.is_file():
        print(f"Missing batch: {batch_path}", file=sys.stderr)
        return 1

    after = load_batch(batch_path)
    wb = load_workbook(excel_path, read_only=True, data_only=True)

    changes: list[dict] = []

    for sn in SHEETS:
        if sn not in wb.sheetnames:
            continue
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = list(rows[0])
        col_idx = {name: header.index(name) for name in TRACK_COLS if name in header}
        if "DevOps link" not in col_idx:
            continue

        for rownum, row in enumerate(rows[1:], start=2):
            row = list(row) + [None] * (len(header) - len(row))
            w = wid_from(row[col_idx["DevOps link"]])
            if not w or w not in after:
                continue
            exp = after[w]
            wid_i = int(w)

            before_status = STATUS_BEFORE.get(wid_i)
            if before_status is not None or wid_i in STATUS_BEFORE:
                cur_status = row[col_idx.get("ADO Status", 0)] if "ADO Status" in col_idx else None
                after_status = exp["state"]
                b = before_status if before_status is not None else cur_status
                if str(b or "").strip() != str(after_status).strip():
                    changes.append(
                        {
                            "id": wid_i,
                            "title": exp["title"][:72],
                            "column": "ADO Status",
                            "before": b,
                            "after": after_status,
                            "where": f"{sn} row {rownum}",
                        }
                    )

            for col, key in [
                ("Work item type", "type"),
                ("DevOps link", "link"),
                ("Sprint", "sprint"),
                ("Tags", "tags"),
                ("Opprettet", "created"),
                ("Sist oppdatert", "changed"),
            ]:
                if col not in col_idx:
                    continue
                cur = row[col_idx[col]]
                if cur is None:
                    cur = ""
                expv = exp[key]
                if wid_i in (445912, 446510) and col == "DevOps link" and not str(cur).strip():
                    before = None
                else:
                    before = cur
                if str(before).strip() != str(expv).strip():
                    if col == "ADO Status":
                        continue
                    changes.append(
                        {
                            "id": wid_i,
                            "title": exp["title"][:72],
                            "column": col,
                            "before": before if before != "" else None,
                            "after": expv,
                            "where": f"{sn} row {rownum}",
                        }
                    )

    wb.close()

    # Dedupe by id + column + before + after
    seen: set[tuple] = set()
    uniq: list[dict] = []
    for c in changes:
        k = (c["id"], c["column"], str(c["before"]), str(c["after"]))
        if k in seen:
            continue
        seen.add(k)
        uniq.append(c)

    uniq.sort(key=lambda x: (x["id"], x["column"]))
    print(json.dumps(uniq, ensure_ascii=False, indent=2))
    print(f"\nTotal distinct cell changes: {len(uniq)}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
