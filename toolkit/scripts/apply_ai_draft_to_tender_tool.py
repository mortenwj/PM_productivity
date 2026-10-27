"""Append AI draft columns to tender compliance Excel tools.

Writes only the AI draft block (columns W onward on Production Plan by default).
Never modifies existing requirement columns unless --promote is explicitly passed
(not implemented in v1 — reserved for future use).

Example:
    uv run python toolkit/scripts/apply_ai_draft_to_tender_tool.py \\
        --drafts docs/tender-reference/ai-drafts/uim-MK-6.json \\
        --req-id MK-6
"""

from __future__ import annotations

import argparse
import json
import sys
from copy import copy
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import column_index_from_string
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SOURCE = REPO_ROOT / "docs/tender-reference/source/UIM ESDH Tender tool.xlsx"
DEFAULT_SHEET = "Production Plan"
LEGEND_SHEET = "AI — Legend"
REQ_ID_COL = "B"
DOCUMENT_COL = "C"
HEADER_ROW = 1
FIRST_AI_COL = "W"

AI_HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9E2F3")
AI_HEADER_FONT = Font(bold=True, color="1F3864")

AI_COLUMNS: list[tuple[str, str, str]] = [
    ("generated", "AI — Generated", "W"),
    ("source", "AI — Source", "X"),
    ("compliance", "AI — Compliance", "Y"),
    ("confidence", "AI — Confidence", "Z"),
    ("summary", "AI — Summary", "AA"),
    ("gaps", "AI — Gaps & clarifications", "AB"),
    ("evidence", "AI — Evidence", "AC"),
    ("next_action", "AI — Next action", "AD"),
    ("reviewed_by", "AI — Reviewed by", "AE"),
    ("review_status", "AI — Review status", "AF"),
    ("run_id", "AI — Run ID", "AG"),
]

LEGEND_ROWS = [
    ("Purpose", "AI-generated draft compliance input. Not official bid text until reviewed."),
    ("Editable columns", "W–AG on Production Plan only. Columns A–V are human/tool fields."),
    ("Promotion", "Copy accepted values to O/P manually after setting Review status = Accepted."),
    ("Regeneration", "Re-running the script overwrites AI columns for matched ReqID rows only."),
    ("Disambiguation", "Use document (Bilag) in draft JSON when ReqID appears on multiple rows."),
]


@dataclass
class DraftEntry:
    req_id: str
    document: str | None = None
    generated: str | None = None
    source: str = "Cursor / tender-response"
    compliance: str = ""
    confidence: str = ""
    summary: str = ""
    gaps: str = ""
    evidence: str = ""
    next_action: str = ""
    reviewed_by: str = ""
    review_status: str = "Draft"
    run_id: str = ""

    @classmethod
    def from_dict(cls, data: dict) -> DraftEntry:
        known = {key for key, _, _ in AI_COLUMNS}
        extra = set(data) - known - {"req_id", "document"}
        if extra:
            raise ValueError(f"Unknown draft fields: {sorted(extra)}")
        return cls(
            req_id=str(data["req_id"]).strip(),
            document=(str(data["document"]).strip() if data.get("document") else None),
            generated=data.get("generated"),
            source=str(data.get("source", "Cursor / tender-response")),
            compliance=str(data.get("compliance", "")),
            confidence=str(data.get("confidence", "")),
            summary=str(data.get("summary", "")),
            gaps=str(data.get("gaps", "")),
            evidence=str(data.get("evidence", "")),
            next_action=str(data.get("next_action", "")),
            reviewed_by=str(data.get("reviewed_by", "")),
            review_status=str(data.get("review_status", "Draft")),
            run_id=str(data.get("run_id", "")),
        )


def load_drafts(path: Path, req_id: str | None) -> list[DraftEntry]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    entries = [DraftEntry.from_dict(item) for item in payload.get("drafts", [])]
    if req_id:
        entries = [e for e in entries if e.req_id == req_id]
    if not entries:
        raise SystemExit(f"No draft entries found{' for ' + req_id if req_id else ''} in {path}")
    return entries


def ai_col_map() -> dict[str, int]:
    return {key: column_index_from_string(col_letter) for key, _, col_letter in AI_COLUMNS}


def ensure_ai_headers(ws: Worksheet) -> dict[str, int]:
    col_by_key = ai_col_map()
    for key, header, col_letter in AI_COLUMNS:
        col_idx = column_index_from_string(col_letter)
        cell = ws.cell(row=HEADER_ROW, column=col_idx, value=header)
        cell.fill = copy(AI_HEADER_FILL)
        cell.font = copy(AI_HEADER_FONT)
    return col_by_key


def cell_text(value: object | None) -> str:
    if value is None:
        return ""
    return str(value).strip()


def find_matching_rows(ws: Worksheet, draft: DraftEntry) -> list[int]:
    matches: list[int] = []
    for row_idx in range(HEADER_ROW + 1, ws.max_row + 1):
        req_val = cell_text(ws.cell(row=row_idx, column=column_index_from_string(REQ_ID_COL)).value)
        if req_val != draft.req_id:
            continue
        if draft.document:
            doc_val = cell_text(ws.cell(row=row_idx, column=column_index_from_string(DOCUMENT_COL)).value)
            if doc_val != draft.document:
                continue
        matches.append(row_idx)
    return matches


def apply_draft_row(ws: Worksheet, row_idx: int, draft: DraftEntry, col_by_key: dict[str, int]) -> None:
    values = {
        "generated": draft.generated or date.today().isoformat(),
        "source": draft.source,
        "compliance": draft.compliance,
        "confidence": draft.confidence,
        "summary": draft.summary,
        "gaps": draft.gaps,
        "evidence": draft.evidence,
        "next_action": draft.next_action,
        "reviewed_by": draft.reviewed_by,
        "review_status": draft.review_status,
        "run_id": draft.run_id,
    }
    for key, col_idx in col_by_key.items():
        ws.cell(row=row_idx, column=col_idx, value=values.get(key, ""))


def ensure_legend_sheet(wb: Workbook) -> None:
    if LEGEND_SHEET in wb.sheetnames:
        return
    ws = wb.create_sheet(LEGEND_SHEET)
    ws["A1"] = "Field"
    ws["B1"] = "Description"
    ws["A1"].font = copy(AI_HEADER_FONT)
    ws["B1"].font = copy(AI_HEADER_FONT)
    ws["A1"].fill = copy(AI_HEADER_FILL)
    ws["B1"].fill = copy(AI_HEADER_FILL)
    for idx, (field, desc) in enumerate(LEGEND_ROWS, start=2):
        ws.cell(row=idx, column=1, value=field)
        ws.cell(row=idx, column=2, value=desc)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 90


def default_output_path(source: Path) -> Path:
    stem = source.stem
    if stem.endswith(" — AI draft"):
        return source
    return source.with_name(f"{stem} — AI draft{source.suffix}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE, help="Input .xlsx")
    parser.add_argument("--output", type=Path, help="Output .xlsx (default: <source> — AI draft.xlsx)")
    parser.add_argument("--drafts", type=Path, required=True, help="JSON file with draft entries")
    parser.add_argument("--sheet", default=DEFAULT_SHEET, help="Worksheet name")
    parser.add_argument("--req-id", help="Apply only this ReqID from the drafts file")
    parser.add_argument("--in-place", action="store_true", help="Overwrite --source instead of writing --output")
    parser.add_argument("--dry-run", action="store_true", help="Print actions without saving")
    parser.add_argument(
        "--promote",
        action="store_true",
        help="Reserved: promote AI values into official columns (not implemented)",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.promote:
        raise SystemExit("--promote is reserved and not implemented. Update O/P manually after review.")

    source = args.source.resolve()
    if not source.exists():
        raise SystemExit(f"Source not found: {source}")

    output = source if args.in_place else (args.output or default_output_path(source)).resolve()
    drafts = load_drafts(args.drafts.resolve(), args.req_id)

    if args.dry_run:
        print(f"Source: {source}")
        print(f"Output: {output}")
        print(f"Sheet: {args.sheet}")
        print(f"AI columns: {FIRST_AI_COL}–{AI_COLUMNS[-1][2]}")
        for draft in drafts:
            doc = f" / {draft.document}" if draft.document else ""
            print(f"  Draft: {draft.req_id}{doc}")
        return 0

    wb = load_workbook(source)
    if args.sheet not in wb.sheetnames:
        raise SystemExit(f"Sheet not found: {args.sheet!r}. Available: {wb.sheetnames}")

    ws = wb[args.sheet]
    col_by_key = ensure_ai_headers(ws)
    ensure_legend_sheet(wb)

    applied: list[str] = []
    skipped: list[str] = []

    for draft in drafts:
        rows = find_matching_rows(ws, draft)
        label = draft.req_id + (f" ({draft.document})" if draft.document else "")
        if not rows:
            skipped.append(label)
            continue
        for row_idx in rows:
            apply_draft_row(ws, row_idx, draft, col_by_key)
            applied.append(f"{label} -> row {row_idx}")

    if not applied:
        raise SystemExit(f"No rows updated. Unmatched drafts: {', '.join(skipped) or 'none'}")

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)

    print(f"Saved: {output}")
    print(f"Updated {len(applied)} row(s):")
    for line in applied:
        print(f"  - {line}")
    if skipped:
        print("Unmatched drafts:")
        for line in skipped:
            print(f"  - {line}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
