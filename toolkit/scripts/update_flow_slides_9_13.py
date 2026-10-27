"""Update Flow section slides 9-13 in 6.12 internal release deck."""
from __future__ import annotations

import shutil
import sys
from datetime import datetime
from pathlib import Path

from pptx import Presentation
PPTX = Path(
    r"C:\Users\MortenJensen\OneDrive - Tieto\Shortcuts\SI Products 360PM - release presentations"
    r"\Internal release presentation public 360 Sept_Okt (6.12).pptx"
)


def hide_slide(prs: Presentation, slide_index: int) -> None:
    sld_id = prs.slides._sldIdLst[slide_index]
    sld_id.set("hidden", "1")


def set_paragraphs(shape, lines: list[str], *, bullet_from: int = 0) -> None:
    tf = shape.text_frame
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        if i >= bullet_from and line.strip():
            p.level = 0


def main() -> int:
    if not PPTX.is_file():
        print(f"Not found: {PPTX}", file=sys.stderr)
        return 1

    backup = PPTX.with_name(
        PPTX.stem + f".backup-flow9-13-{datetime.now().strftime('%Y%m%d-%H%M%S')}" + PPTX.suffix
    )
    shutil.copy2(PPTX, backup)
    print(f"Backup: {backup}")

    prs = Presentation(str(PPTX))

    # Slide 9 — section title + GA line
    s9 = prs.slides[8]
    s9.shapes[2].text_frame.paragraphs[0].text = "Better flow. More value"
    set_paragraphs(
        s9.shapes[3],
        [
            "360° Flow",
            "Generally available · Public 360° 6.12+",
        ],
    )

    # Slide 10 — hide (compliance one-liner duplicated elsewhere / webinar)
    hide_slide(prs, 9)

    # Slide 11 — trimmed “what is Flow”
    s11 = prs.slides[10]
    set_paragraphs(
        s11.shapes[3],
        [
            "360° Flow adds structured, process-oriented case handling in ProcessView.",
            "Process models, phases, activities, checklists, and milestones — clear next steps and a common standard.",
            "Case handlers: guided work in the phase; managers: status, deadlines, and workload across cases.",
            "Organizations: standardize processes while keeping flexibility — configured with Tieto or an implementation partner.",
        ],
        bullet_from=1,
    )

    # Slide 12 — hide (merged into slide 11)
    hide_slide(prs, 11)

    # Slide 13 — 6.12 highlights
    s13 = prs.slides[12]
    set_paragraphs(
        s13.shapes[3],
        [
            "Flow platform — highlights in 6.12",
            "Background processes on a stage (Process Model Admin); workers see status and can retry — toggle FT:EnableBackgroundProcesses when agreed.",
            "Email and meetings from processes — reply-with-mail and Outlook appointments when steps run immediately.",
            "Process Model Admin — drag-and-drop reorder processes, milestones, and stage templates.",
            "Process email dialog — standard templates and macros; optional template and language.",
            "Advanced search — filters on stage, milestone, progress plan, theme, and related criteria (all customers).",
            "Deadlines — optional no-deadline calculation; UI shows “No deadline” where applicable.",
        ],
        bullet_from=1,
    )

    prs.save(str(PPTX))
    print(f"Updated: {PPTX}")
    print("Slides 9, 11, 13 updated; slides 10 and 12 hidden; slide 14 unchanged.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
