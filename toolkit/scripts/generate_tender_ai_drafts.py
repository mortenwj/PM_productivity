"""Generate AI draft JSON for all rows in a tender compliance workbook.

Reads Production Plan, maps tool ratings to AI compliance fields, enriches with
tender-reference evidence grep and optional per-requirement overrides.

Example:
    uv run python toolkit/scripts/generate_tender_ai_drafts.py \\
        --source docs/tender-reference/source/UIM\\ ESDH\\ Tender\\ tool.xlsx \\
        --output docs/tender-reference/ai-drafts/uim-full.json \\
        --overrides docs/tender-reference/ai-drafts/uim-MK-6.json
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import asdict, dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SOURCE = REPO_ROOT / "docs/tender-reference/source/UIM ESDH Tender tool.xlsx"
DEFAULT_OUTPUT = REPO_ROOT / "docs/tender-reference/ai-drafts/uim-full.json"
DEFAULT_SHEET = "Production Plan"
RED_FLAGS_SHEET = "Red flags"

EVIDENCE_DIRS = [
    REPO_ROOT / "docs/tender-reference/extracted",
    REPO_ROOT / "docs/product",
]

STOPWORDS = {
    "og", "i", "på", "af", "til", "en", "et", "der", "som", "skal", "kan", "med",
    "for", "at", "den", "det", "de", "eller", "ikke", "fra", "ved", "være", "have",
    "the", "and", "for", "with", "that", "shall", "must", "support", "system",
    "esdh", "løsningen", "løsningens", "brugeren", "bruger", "kunden", "leverandøren",
}

TOOL_TO_COMPLIANCE = {
    "1. Fully met": "Met",
    "2. Partly met": "Partial",
    "3. Not able to meet": "Not met",
}

COL = {
    "req_id": "B",
    "document": "C",
    "section": "D",
    "title": "E",
    "requirement": "F",
    "type": "G",
    "action": "H",
    "can_meet": "O",
    "comment": "P",
    "responsible": "L",
}


@dataclass
class RequirementRow:
    row_idx: int
    req_id: str
    document: str
    section: str
    title: str
    requirement: str
    req_type: str
    action: str
    can_meet: str
    comment: str
    responsible: str


@dataclass
class DraftOut:
    req_id: str
    document: str
    generated: str
    source: str
    compliance: str
    confidence: str
    summary: str
    gaps: str
    evidence: str
    next_action: str
    reviewed_by: str
    review_status: str
    run_id: str
    _row_idx: int = field(default=0, repr=False)


def cell_val(ws, row: int, col_letter: str) -> str:
    value = ws.cell(row=row, column=column_index_from_string(col_letter)).value
    if value is None:
        return ""
    return str(value).strip()


def load_rows(path: Path, sheet: str) -> list[RequirementRow]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise SystemExit(f"Sheet not found: {sheet!r}")
    ws = wb[sheet]
    rows: list[RequirementRow] = []
    for row_idx in range(2, ws.max_row + 1):
        req_id = cell_val(ws, row_idx, COL["req_id"])
        if not req_id or req_id == "ReqID":
            continue
        rows.append(
            RequirementRow(
                row_idx=row_idx,
                req_id=req_id,
                document=cell_val(ws, row_idx, COL["document"]),
                section=cell_val(ws, row_idx, COL["section"]),
                title=cell_val(ws, row_idx, COL["title"]),
                requirement=cell_val(ws, row_idx, COL["requirement"]),
                req_type=cell_val(ws, row_idx, COL["type"]),
                action=cell_val(ws, row_idx, COL["action"]),
                can_meet=cell_val(ws, row_idx, COL["can_meet"]),
                comment=cell_val(ws, row_idx, COL["comment"]),
                responsible=cell_val(ws, row_idx, COL["responsible"]),
            )
        )
    wb.close()
    return rows


def load_red_flags(path: Path) -> dict[str, str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if RED_FLAGS_SHEET not in wb.sheetnames:
        wb.close()
        return {}
    ws = wb[RED_FLAGS_SHEET]
    flags: dict[str, str] = {}
    for row_idx in range(2, ws.max_row + 1):
        req_id = cell_val(ws, row_idx, "A")
        comment = cell_val(ws, row_idx, "C")
        if req_id and comment:
            flags[req_id] = comment
    wb.close()
    return flags


def load_overrides(paths: list[Path]) -> dict[tuple[str, str], dict]:
    merged: dict[tuple[str, str], dict] = {}
    for path in paths:
        if not path.exists():
            continue
        payload = json.loads(path.read_text(encoding="utf-8"))
        for item in payload.get("drafts", []):
            key = (str(item["req_id"]).strip(), str(item.get("document", "")).strip())
            merged[key] = item
    return merged


def tokenize(text: str) -> set[str]:
    words = re.findall(r"[a-zA-ZæøåÆØÅ0-9]{4,}", text.lower())
    return {w for w in words if w not in STOPWORDS}


def search_evidence(row: RequirementRow, max_hits: int = 3) -> tuple[str, str]:
    terms = tokenize(f"{row.title} {row.requirement[:400]}")
    if not terms:
        return "Inferred", "Ingen nøgleord til evidens-søgning."

    scored: list[tuple[int, str]] = []
    for base in EVIDENCE_DIRS:
        if not base.exists():
            continue
        for path in base.rglob("*"):
            if path.suffix.lower() not in {".txt", ".md"}:
                continue
            if path.name.startswith("_"):
                continue
            try:
                text = path.read_text(encoding="utf-8", errors="ignore").lower()
            except OSError:
                continue
            score = sum(1 for t in terms if t in text)
            if score >= 2:
                rel = path.relative_to(REPO_ROOT).as_posix()
                scored.append((score, rel))

    if not scored:
        return "Lead", "Ingen prior-art træffere i docs/tender-reference/extracted eller docs/product."

    scored.sort(key=lambda x: (-x[0], x[1]))
    unique: list[str] = []
    for _, rel in scored:
        if rel not in unique:
            unique.append(rel)
        if len(unique) >= max_hits:
            break
    confidence = "Verified" if scored[0][0] >= 4 else "Lead"
    return confidence, "; ".join(unique)


def compliance_for_row(row: RequirementRow) -> str:
    if row.can_meet in TOOL_TO_COMPLIANCE:
        return TOOL_TO_COMPLIANCE[row.can_meet]
    action = row.action.lower()
    if action == "answer":
        return "Draft needed"
    if action == "check":
        return "Clarify"
    return "Clarify"


def confidence_for_row(row: RequirementRow, evidence_confidence: str) -> str:
    if row.can_meet == "1. Fully met":
        return "Verified" if evidence_confidence == "Verified" else "Lead"
    if row.can_meet == "3. Not able to meet":
        return "Verified" if row.comment else "Lead"
    if row.can_meet == "2. Partly met":
        return "Lead"
    return evidence_confidence


def summary_for_row(row: RequirementRow, compliance: str) -> str:
    action = row.action or "Verify"
    title = row.title or row.req_id
    if compliance == "Met":
        return f"{action}: {title} — vurderet opfyldt (tool: {row.can_meet or 'n/a'})."
    if compliance == "Partial":
        return f"{action}: {title} — delvist opfyldt; afklar scope og scoring."
    if compliance == "Not met":
        return f"{action}: {title} — gap ift. kravtekst; eskalér commitment eller afklaring."
    if compliance == "Draft needed":
        return f"{action}: {title} — informationskrav; udkast bilag/dokumentation mangler."
    return f"{action}: {title} — kræver produkt-/team-afklaring (ingen tool-rating)."


def gaps_for_row(row: RequirementRow, red_flags: dict[str, str]) -> str:
    parts: list[str] = []
    if row.comment:
        parts.append(row.comment)
    flag = red_flags.get(row.req_id)
    if flag and flag not in parts:
        parts.append(f"Red flag: {flag}")
    if row.can_meet == "3. Not able to meet" and not parts:
        parts.append("Tool markeret som ikke opfyldt — bekræft gap og mitigation.")
    if not row.can_meet and row.action.lower() == "answer":
        parts.append("Mangler link til svar-bilag i kolonne T.")
    if not parts:
        return ""
    return " | ".join(parts)


def next_action_for_row(row: RequirementRow, compliance: str) -> str:
    owner = row.responsible or "Owner"
    if compliance == "Met":
        return f"{owner}: Confirm verify-tekst; opdatér review status efter tjek."
    if compliance == "Not met":
        return f"{owner}: Commitment/no-bid beslutning; overvej Q&A til ordregiver."
    if compliance == "Draft needed":
        return f"{owner}: Opret/link svar-bilag; draft IK-tekst fra prior art."
    if compliance == "Partial":
        return f"{owner}: Afklar delkrav; align tool-rating (O) med AI compliance."
    return f"{owner}: Udfyld tool-rating (O) og evidens; gennemgå AI-kladde."


def apply_mcp_overlay(draft: DraftOut, mcp: dict) -> None:
    if mcp.get("ratatoskr") or mcp.get("wiki"):
        parts: list[str] = []
        if mcp.get("ratatoskr"):
            parts.append("Ratatoskr: " + "; ".join(mcp["ratatoskr"][:3]))
        if mcp.get("wiki"):
            parts.append("Wiki: " + "; ".join(mcp["wiki"][:2]))
        draft.evidence = " | ".join(parts)
    elif mcp.get("evidence_note"):
        draft.evidence = str(mcp["evidence_note"])
    if mcp.get("confidence"):
        draft.confidence = str(mcp["confidence"])
    draft.source = "Cursor / tender-response (MCP-enriched)"


def draft_from_row(
    row: RequirementRow,
    red_flags: dict[str, str],
    overrides: dict[tuple[str, str], dict],
    mcp_cache: dict[tuple[str, str], dict],
    run_id: str,
    generated: str,
    source: str,
) -> DraftOut:
    key = (row.req_id, row.document)
    if key in overrides:
        o = overrides[key]
        return DraftOut(
            req_id=row.req_id,
            document=row.document,
            generated=str(o.get("generated", generated)),
            source=str(o.get("source", source)),
            compliance=str(o.get("compliance", "")),
            confidence=str(o.get("confidence", "")),
            summary=str(o.get("summary", "")),
            gaps=str(o.get("gaps", "")),
            evidence=str(o.get("evidence", "")),
            next_action=str(o.get("next_action", "")),
            reviewed_by=str(o.get("reviewed_by", "")),
            review_status=str(o.get("review_status", "Draft")),
            run_id=str(o.get("run_id", run_id)),
            _row_idx=row.row_idx,
        )

    ev_conf, evidence = search_evidence(row)
    compliance = compliance_for_row(row)
    draft = DraftOut(
        req_id=row.req_id,
        document=row.document,
        generated=generated,
        source=source,
        compliance=compliance,
        confidence=confidence_for_row(row, ev_conf),
        summary=summary_for_row(row, compliance),
        gaps=gaps_for_row(row, red_flags),
        evidence=evidence,
        next_action=next_action_for_row(row, compliance),
        reviewed_by="",
        review_status="Draft",
        run_id=run_id,
        _row_idx=row.row_idx,
    )
    if key in mcp_cache:
        apply_mcp_overlay(draft, mcp_cache[key])
    return draft


def draft_to_json(d: DraftOut) -> dict:
    return {
        "req_id": d.req_id,
        "document": d.document,
        "generated": d.generated,
        "source": d.source,
        "compliance": d.compliance,
        "confidence": d.confidence,
        "summary": d.summary,
        "gaps": d.gaps,
        "evidence": d.evidence,
        "next_action": d.next_action,
        "reviewed_by": d.reviewed_by,
        "review_status": d.review_status,
        "run_id": d.run_id,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--sheet", default=DEFAULT_SHEET)
    parser.add_argument("--overrides", type=Path, action="append", default=[])
    parser.add_argument("--tender", default="UIM ESDH")
    parser.add_argument("--source-label", default="Cursor / tender-response (batch)")
    parser.add_argument("--run-id", default="")
    parser.add_argument(
        "--mcp-cache",
        type=Path,
        default=REPO_ROOT / "docs/tender-reference/ai-drafts/uim-mcp-evidence-cache.json",
        help="MCP evidence cache JSON (0 = disable)",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    source = args.source.resolve()
    if not source.exists():
        raise SystemExit(f"Source not found: {source}")

    generated = date.today().isoformat()
    run_id = args.run_id or f"{args.tender.lower().replace(' ', '-')}-full-{generated}"

    rows = load_rows(source, args.sheet)
    red_flags = load_red_flags(source)
    override_paths = [Path(p).resolve() for p in args.overrides]
    override_paths.append(REPO_ROOT / "docs/tender-reference/ai-drafts/uim-MK-6.json")
    overrides = load_overrides(override_paths)

    mcp_cache: dict[tuple[str, str], dict] = {}
    if args.mcp_cache and str(args.mcp_cache) != "0" and args.mcp_cache.exists():
        mcp_payload = json.loads(args.mcp_cache.read_text(encoding="utf-8"))
        for k, v in mcp_payload.get("entries", {}).items():
            if "|" in k:
                req_id, document = k.split("|", 1)
                mcp_cache[(req_id, document)] = v

    drafts = [
        draft_from_row(row, red_flags, overrides, mcp_cache, run_id, generated, args.source_label)
        for row in rows
    ]

    payload = {
        "tender": args.tender,
        "generated": generated,
        "run_id": run_id,
        "source_workbook": str(source.relative_to(REPO_ROOT)).replace("\\", "/"),
        "row_count": len(drafts),
        "drafts": [draft_to_json(d) for d in drafts],
    }

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    from collections import Counter

    counts = Counter(d.compliance for d in drafts)
    print(f"Wrote {len(drafts)} drafts -> {args.output}")
    print("Compliance:", dict(counts))
    print(f"Overrides applied: {len(overrides)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
