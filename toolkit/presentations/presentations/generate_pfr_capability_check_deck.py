"""Generate a short value-focused deck for the PFR AI capability-check skill."""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt


def _bullet(tf, lines: list[str]) -> None:
    p = tf.add_paragraph()
    p.text = lines[0]
    p.level = 0
    p.font.size = Pt(20)
    for line in lines[1:]:
        p = tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(20)


def main() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide 1 — Title
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "PFR — AI-assisted capability check"
    st = slide.placeholders[1]
    st.text = (
        "Structured first pass on Product First requests:\n"
        "value, evidence, and confidence — before you commit people."
    )
    for p in slide.shapes.title.text_frame.paragraphs:
        p.font.size = Pt(40)
    for p in st.text_frame.paragraphs:
        p.font.size = Pt(22)

    # Slide 2 — Why / value
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Why this exists"
    tf = slide.placeholders[1].text_frame
    tf.clear()
    _bullet(
        tf,
        [
            "Many requests sound like new product gaps when the answer is often "
            "already supported, configurable, or a small change — but proving that takes time.",
            "Teams pay in slow triage, backlog noise, and repeated “do we have this?” debates.",
            "This skill gives a structured first answer: what the request really needs, "
            "what we can show today, and how confident we are — before PM and engineering "
            "carry the full cost.",
        ],
    )
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(36)

    # Slide 3 — Logic + human gate
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "How it thinks (and where humans stay)"
    tf = slide.placeholders[1].text_frame
    tf.clear()
    _bullet(
        tf,
        [
            "Reframe the Feature as a neutral required capability — not only the customer’s frustration.",
            "Compare that to observed capability from verified sources, and separate "
            "configuration or filters from enforced access or guarantees.",
            "Land in one outcome (e.g. already supported, partial gap, true gap, insufficient information) "
            "with explicit unknowns when the request is thin.",
            "Human review is required when confidence is lower or evidence is thin — "
            "the workflow accelerates triage; it does not replace PM or architecture sign-off.",
        ],
    )
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(34)

    out = Path(__file__).resolve().parent / "PFR_AI_capability_check.pptx"
    prs.save(out)
    return out


if __name__ == "__main__":
    path = main()
    print(path)
